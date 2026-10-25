"""Loss-aware PPEF reader. No formula execution, no silent column dropping, no name splitting."""
import csv,hashlib,io,json,re,unicodedata,zipfile
from datetime import date,datetime
from decimal import Decimal,InvalidOperation
from openpyxl import load_workbook
from .registry_models import Register,Money,validate_dates,quality,CATALOGUES


def norm(s):return re.sub(r'[^a-z0-9]+','',unicodedata.normalize('NFKD',str(s or '')).encode('ascii','ignore').decode().lower())
# Input columns have a canonical destination. Spreadsheet calculations are retained in source only.
INPUT={
 'rp_number':['N° RP','rp_number'],'date_rp':['Date RP','date_rp'],'rp_year':['rp_year'],
 'provenance':['Unité (OPJ) de Provenance','Provenance','provenance'],'transnational':['Dossier à caractère transnational (oui/non)','transnational'],
 'names':['Nom du prévenu','Inculpé'],'person_physical_count':['Nombre prévenus (PP)','person_physical_count'],'person_legal_count':['Nombre prévenus (PM)','person_legal_count'],
 'party_type':['Type de prévenus',"Type d’ inculpés","Type d' inculpés"],'profession':['profession (personne publique )'],'activity':["Domaine d' activité (personne morale)","Domaine d' activité"],
 'category':['Infraction retenue par le Parquet','Infraction','category'],'offence_amount':['montant infraction'],'orientation':['Suite donnée'],'legal_form':['Forme juridique'],
 'ri_number':['N° RI','ri_number'],'cabinet':['Cabinet'],'closure_date':['Date de clôture ordonnance'],'closure_kind':["Nature de l'ordonnance de clôture"],
 'seizure_date':['Date de la saisie'],'seizure_object':['Objet Saisi','Saisie'],'seizure_amount':['Montant saisie'],
 'first_hearing':['Date de la 1ère audience'],'tribunal_label':['Infraction retenue par le tribunal','Nature infraction retenue par le tribunal'],
 'decision_number':['N° de la décision'],'decision_date':['Date de la décision'],'decision_nature':['Caractère décision au fond'],
 'decision_content':['Contenu de la décision'],'sentence':['Peines prononcées','Peine'],'fine':['Amende en fcfa','Amende'],
 'confiscation_object':['Objet confisqué','Confiscation'],'confiscation_amount':['Montant confisqué'],'trafficking':['Est-ce un cas de traite des personnes (OUI/NON)'],
 'convicted_total':['Total condamné'],'convicted_male':['Dont condamné masculin'],'convicted_female':['Dont condamné féminin'],
 'convicted_nationalities':['Nationalité des condamnés (Préciser nombre condamnés par nationalité)'],'victim_total':['Total victimes'],
 'position':['POSITION','Position'],'observations':['Observations'],'harmonised_label':['Harmoisation des infractions'],
 'compensation':['Réparation accordée aux victimes (en préciser le montant par victime)'],'jurisdiction':['Juridiction ayant statué'],
 'source_date':['Date'],'order_number':["N° D'Ordre"],'specific_offence':['Infraction spécifique'],'other_specificity':['Autre spécificité'],'offence_treatment':['Traitement infraction'],'matter':['Matière']}
DERIVED=['Compte infraction','Corruption','Infractions assimilées','Abus de fonction','Abus de biens sociaux','Concussion','Extorsion de fonds','Détournement de deniers publics','Blanchiment autonome','fraude fiscale','crimes environnementaux','Cybercriminalité','Trafic de Stupéfiants','Traite des personnes','Financement du terrorisme','Commercialisation illicite de produits pharmaceutiques','Commercialisation illicite de produits agricoles','contenu de la decision 1','Condamné','Relaxé',"nature de l'ordonnance de cloture",'ONL','JUGE','juge2','Type de prevenu 2','Personne morale','Personne physique','Personne publique','confis','confisqué','decompte confiscation','Objets saisies','saisie','decompte saisie','Dossiers à caractère transnational','decompte dossiers',"Durée entre 1ère évocation et la décision","Durée entre 1ère évocation et transmission au GC",'Mois transmission GC','Mois 1ère évocation','Mois décision','Année transmission GC','Année 1ère évocation','Année décision','Evoqué avant 2022','Vidé avant 2022','Transmis GC avant 2022','Colonne12']
ALIASES={norm(label):key for key,labels in INPUT.items() for label in labels}
ALL_DERIVED_KEYS={norm(x) for x in DERIVED}
DERIVED_KEYS=ALL_DERIVED_KEYS-set(ALIASES)

def plain(v):
 if isinstance(v,(datetime,date)):return v.isoformat()
 return v

def date_value(v):
 if isinstance(v,datetime):return v.date().isoformat()
 if isinstance(v,date):return v.isoformat()
 if v in (None,''):return None
 for fmt in ('%Y-%m-%d','%Y-%m-%dT%H:%M:%S','%d/%m/%Y','%d-%m-%Y','%Y/%m/%d'):
  try:return datetime.strptime(str(v).strip(),fmt).date().isoformat()
  except ValueError:pass
 raise ValueError('Date invalide : format attendu JJ/MM/AAAA ou AAAA-MM-JJ')

def integer(v):
 if v is None or str(v).strip()=='':return None
 try:n=Decimal(str(v).replace(' ',''))
 except InvalidOperation:raise ValueError('Effectif entier requis')
 if not n.is_finite() or not 0<=n<=2147483647 or n!=n.to_integral_value():raise ValueError('Effectif entier positif ou nul requis')
 return int(n)

def money(v):
 if v is None or str(v).strip()=='':return Money().model_dump()
 text=str(v).strip().replace('\u00a0','').replace(' ','').replace(',','.')
 text=re.sub(r'(?i)(fcfa|xof)$','',text)
 if norm(text) in ('nd','nondetermine','inconnu','aevaluer'):return Money(state='undetermined').model_dump()
 # Floats from Excel are parsed as decimal text, not truncated or silently coerced.
 try:n=Decimal(text)
 except InvalidOperation:raise ValueError('Montant non numérique')
 if not n.is_finite() or n<0 or n.adjusted()>39 or n.as_tuple().exponent < -2:raise ValueError('Montant invalide ou précision supérieure à 2 décimales')
 return Money(state='known',value=format(n,'f')).model_dump()

def choice(value,key):
 from .registry_models import CATALOGUES
 if value is None or str(value).strip()=='':return {key:'not_recorded'}
 for code,label in CATALOGUES[key].items():
  if norm(value) in (norm(code),norm(label)):return {key:code}
 return {key:'other',key+'_other':str(value)}

def convert(raw,line):
 def t(key):return str(raw[key]).strip() if raw.get(key) not in (None,'') else None
 received=date_value(raw.get('date_rp'))
 # Feuil1's generic Date is retained, never silently relabelled as Date RP.
 if not t('rp_number') or not received:raise ValueError('Numéro RP et Date RP requis. Une date générique de Feuil1 ne peut être assimilée automatiquement à la Date RP.')
 year=integer(raw.get('rp_year')) or int(received[:4])
 if year!=int(received[:4]):raise ValueError('Année RP différente de la date RP')
 trans=norm(raw.get('transnational'));trans_state={'oui':'yes','yes':'yes','true':'yes','1':'yes','non':'no','no':'no','false':'no','0':'no'}.get(trans,'undetermined')
 trans_choice=choice(raw.get('transnational'),'transnational_state')
 if trans in ('oui','yes','true','1','non','no','false','0'):trans_choice={'transnational_state':trans_state}
 trans_state=trans_choice['transnational_state']
 core={'rp_number':t('rp_number'),'date_rp':received,'rp_year':year,'provenance':t('provenance') or 'Non renseigné','category':t('category') or 'Qualification non déterminée','ri_number':t('ri_number'),'transnational':trans_state=='yes','person_physical_count':integer(raw.get('person_physical_count')),'person_legal_count':integer(raw.get('person_legal_count'))}
 procedure={'cabinet':t('cabinet'),'closure_date':date_value(raw.get('closure_date')),'first_hearing_planned':None,'position':t('position'),'observations':t('observations'),**choice(raw.get('orientation'),'orientation'),**choice(raw.get('closure_kind'),'closure_kind')}
 record={'ri_state':'assigned' if t('ri_number') else 'not_recorded','source_context':{**{k:t(k) for k in ('source_date','order_number','specific_offence','other_specificity','offence_treatment','matter')},'first_hearing_source_date':date_value(raw.get('first_hearing')),'party_type_source':t('party_type'),'trafficking_source_text':t('trafficking')},'jurisdiction_name':t('jurisdiction'),**trans_choice,'procedure':procedure,'count_basis':'declared','people':[],
 'declarations':{k:(integer(raw.get(k)) if k in ('convicted_total','convicted_male','convicted_female','victim_total') else t(k)) for k in ('convicted_total','convicted_male','convicted_female','victim_total','convicted_nationalities')},
 'offences':[],'decisions':[],'sentences':[],'measures':[]}
 if raw.get('names'):
  pp=core['person_physical_count'];pm=core['person_legal_count']
  if (pp,pm) in ((1,0),(0,1)):
   record['people']=[{'reference':'P'+str(line),'display_name':t('names'),'kind':'physical' if pp==1 else 'legal',**choice(raw.get('profession'),'profession'),**choice(raw.get('activity'),'activity'),**choice(raw.get('legal_form'),'legal_form')}]
   if pp==1:record['people'][0].update(legal_form='not_applicable',activity='not_applicable')
  else:record['declarations']['unresolved_people_text']=t('names')
 if t('category') or t('tribunal_label') or t('offence_amount'):
  record['offences']=[{'reference':'I'+str(line),'parquet_label':t('category'),'tribunal_label':t('tribunal_label'),'harmonised_label':t('harmonised_label'),'amount':money(raw.get('offence_amount'))}]
 decision=None
 if any(t(k) for k in ('decision_number','decision_date','decision_content','decision_nature','sentence')):
  decision={'reference':'D'+str(line),'number':t('decision_number'),'decision_date':date_value(raw.get('decision_date')),'content':t('decision_content'),**choice(raw.get('decision_nature'),'nature')}
  result_key=lambda value:re.sub(r'\s+',' ',unicodedata.normalize('NFKD',str(value or '')).encode('ascii','ignore').decode().casefold()).strip()
  content=result_key(raw.get('decision_content'))
  # A word inside a narrative is not a judicial result (e.g. annulment or negation).
  exact_results={result_key(value):code for code,label in CATALOGUES['result'].items() for value in (code,label)}
  exact_results.update({result_key(value):code for value,code in [('Condamné','conviction'),('Relaxe','acquittal'),('Relaxé','acquittal'),('Acquittement','acquittal'),('Condamnation et relaxe','mixed')]})
  decision['result']=exact_results.get(content,'undetermined') if content else 'not_recorded'
  if decision['result']=='other':decision['result_other']=t('decision_content')
  if norm(raw.get('decision_nature')) in ('oui','1','true'):decision.update(nature='merits',nature_other=None)
  if norm(raw.get('decision_nature')) in ('non','0','false'):decision.update(nature='not_merits',nature_other=None)
  record['decisions']=[decision]
 if t('sentence') and decision:record['sentences']=[{'reference':'S'+str(line),'decision_reference':decision['reference'],'sentence_kind':'other','sentence_kind_other':t('sentence'),'source_text':t('sentence')}]
 for kind,date_key,obj_key,amount_key in [('seizure','seizure_date','seizure_object','seizure_amount'),('confiscation',None,'confiscation_object','confiscation_amount'),('fine',None,None,'fine'),('compensation',None,None,'compensation')]:
  if any(t(k) for k in (date_key,obj_key,amount_key) if k):
   m={'reference':kind+str(line),'kind':kind,'measure_date':date_value(raw.get(date_key)) if date_key else (decision['decision_date'] if decision else None),'amount':money(raw.get(amount_key)),'description':t(obj_key) if obj_key else None}
   if obj_key:m.update(choice(raw.get(obj_key),'object_kind'))
   if decision and kind!='seizure':m['decision_reference']=decision['reference']
   record['measures'].append(m)
 if norm(raw.get('trafficking')) in ('oui','non'):record['trafficking_state']='yes' if norm(raw['trafficking'])=='oui' else 'no'
 result=validate_dates(Register.model_validate(record),date.fromisoformat(received))
 return {'core':core,'record':result.model_dump(mode='json'),'source_rows':[line]}


def preview(filename,content,sheet=None,date_role=None):
 if filename.lower().endswith('.jsonl'):return preview_package(content)
 coverage=[];warnings=[];errors=[];sheets=[];source_rows=[]
 if filename.lower().endswith('.csv'):
  reader=csv.reader(io.StringIO(content.decode('utf-8-sig')));rows=list(reader)
  if any(len(row)>256 for row in rows):raise ValueError('Maximum 256 colonnes')
  if len(rows)>10001:raise ValueError('Maximum 10 000 lignes')
  worksheets=[('CSV',rows)]
 elif filename.lower().endswith('.xlsx'):
  with zipfile.ZipFile(io.BytesIO(content)) as z:
   if sum(x.file_size for x in z.infolist())>60*1024*1024:raise ValueError('Classeur décompressé limité à 60 Mo')
  w=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
  worksheets=[]
  try:
   for ws in w:
    if (ws.max_row or 0)>10004 or (ws.max_column or 0)>256:raise ValueError('Maximum 10 000 lignes et 256 colonnes par feuille')
    sheets.append({'name':ws.title,'rows':ws.max_row,'columns':ws.max_column,'role':'source' if ws.title in ('AUDIENCES PENALES','Feuil1') else 'reference_or_calculation'})
    if ws.title in ('AUDIENCES PENALES','Feuil1') or (sheet and ws.title==sheet):worksheets.append((ws.title,list(ws.iter_rows(values_only=True))))
  finally:w.close()
 else:raise ValueError('Formats acceptés : CSV et XLSX')
 candidates=[]
 for title,rows in worksheets:
  for idx,head in enumerate(rows[:6]):
   if any(ALIASES.get(norm(x))=='rp_number' for x in head):
    from openpyxl.utils import get_column_letter
    # CSV rows may be wider than their header. Preserve and reject those extra cells.
    width=max((len(row) for row in rows[idx:]),default=len(head))
    head=list(head)+[None]*(width-len(head))
    labels=[];seen_headers=set()
    for i,x in enumerate(head):
     label=str(x).strip() if x is not None and str(x).strip() else 'COL_'+str(i+1)
     if norm(label) in seen_headers:
      if title!='AUDIENCES PENALES':errors.append({'row':idx+1,'errors':['En-têtes dupliqués : '+title]})
      label+=' ['+get_column_letter(i+1)+']'
     seen_headers.add(norm(label));labels.append(label)
    extracted=[]
    for line,values in enumerate(rows[idx+1:],idx+2):
     # Formula-only rows in report templates are not records.
     if not any(v is not None and str(v).strip() and not (isinstance(v,str) and v.startswith('=')) for v in values):continue
     business=[v for col,(h,v) in enumerate(zip(labels,values)) if ALIASES.get(norm(h)) not in ('matter','order_number') and not (norm(h) in DERIVED_KEYS or (title=='AUDIENCES PENALES' and col>=40 and norm(re.sub(r' \[[A-Z]+\]$','',h)) in ALL_DERIVED_KEYS)) and v is not None and str(v).strip() and not (isinstance(v,str) and v.startswith('='))]
     if not business:continue
     extracted.append((line,dict(zip(labels,values))))
    candidates.append((title,labels,extracted));break
 if sheet:
  selected=next((x for x in candidates if x[0]==sheet),None)
  if not selected:raise ValueError('Feuille source ou en-têtes RP introuvables')
 else:
  filled=[x for x in candidates if x[2]]
  if len(filled)>1:raise ValueError('Plusieurs feuilles source renseignées : sélectionner explicitement une feuille pour éviter les doubles comptes')
  selected=filled[0] if filled else candidates[0] if candidates else None
 if not selected:raise ValueError('Aucune feuille source reconnue')
 title,headers,rows=selected
 bindings={}
 for i,h in enumerate(headers):
  k=norm(h)
  derived=(title=='AUDIENCES PENALES' and i>=40 and norm(re.sub(r' \[[A-Z]+\]$','',h)) in ALL_DERIVED_KEYS) or k in DERIVED_KEYS or (h.startswith('COL_') and not any(values.get(h) is not None and str(values[h]).strip() and not (isinstance(values[h],str) and values[h].startswith('=')) for _,values in rows))
  bindings[h]='@derived' if derived else ALIASES.get(k)
  coverage.append({'column':h,'destination':bindings[h] if bindings[h]!='@derived' else None,'handling':'derived_preserved' if derived else 'mapped' if bindings[h] else 'unmapped'})
 groups={}
 for line,values in rows:
  raw={};row_errors=[]
  for h,v in values.items():
   k=norm(h);destination=bindings[h]
   if v in (None,''):continue
   if isinstance(v,str) and v.startswith('='):
    if destination and destination!='@derived':row_errors.append('Champ métier calculé sans valeur source indépendante : '+h)
    continue
   if destination and destination!='@derived':
    if destination in raw and raw[destination]!=v:row_errors.append('Colonnes contradictoires pour '+destination)
    raw[destination]=v
   elif destination!='@derived':row_errors.append('Colonne renseignée sans correspondance : '+h)
  source_rows.append({'sheet':title,'row':line,'cells':{h:plain(v) for h,v in values.items() if v is not None}})
  if row_errors:errors.append({'row':line,'errors':row_errors});continue
  try:
   if title=='Feuil1' and date_role=='rp' and not raw.get('date_rp'):
    raw['date_rp']=raw.get('source_date')
   item=convert(raw,line);key=(item['core']['rp_number'].casefold(),item['core']['rp_year'])
   if key in groups:
    existing=groups[key]
    for part in ('procedure','declarations','jurisdiction_name','source_context'):
     if existing['record'][part]!=item['record'][part]:raise ValueError('Même dossier avec '+part+' contradictoire : rapprochement manuel requis')
    if existing['core']!=item['core']:raise ValueError('Même RP/année avec informations de dossier contradictoires : rapprochement manuel requis')
    # Exact repeated snapshots are not automatically counted twice.
    previous_records=existing.setdefault('_snapshots',[])
    snap=json.dumps({k:v for k,v in raw.items()},sort_keys=True,default=plain)
    if snap in previous_records:warnings.append({'row':line,'message':'Ligne identique conservée dans la source, non recomptée'})
    else:
     measure_keys=lambda items:{json.dumps({k:v for k,v in x.items() if k!='reference'},sort_keys=True) for x in items}
     if measure_keys(existing['record']['measures']) & measure_keys(item['record']['measures']):raise ValueError('Mesure financière répétée : préciser s’il s’agit du même bien ou de deux mesures avant reprise')
     for key2 in ('people','offences','decisions','sentences','measures'):
      if key2=='offences':
       keys=lambda arr:{json.dumps({k:v for k,v in x.items() if k!='reference'},sort_keys=True) for x in arr}
       existing_keys=keys(existing['record'][key2])
       existing['record'][key2]+=[x for x in item['record'][key2] if json.dumps({k:v for k,v in x.items() if k!='reference'},sort_keys=True) not in existing_keys]
      else:existing['record'][key2]+=item['record'][key2]
     previous_records.append(snap)
    existing['source_rows'].append(line)
   else:
    item['_snapshots']=[json.dumps(raw,sort_keys=True,default=plain)];groups[key]=item
  except (ValueError,TypeError) as e:errors.append({'row':line,'errors':[str(e)]})
 items=list(groups.values())
 for item in items:
  snapshots=item.pop('_snapshots',[])
  # If grouped identities may represent the same person across decisions, preserve as review input.
  if len(snapshots)>1 and item['record']['people']:
   text=' | '.join(p['display_name'] or 'Identité inconnue' for p in item['record']['people'])
   item['record']['declarations']['unresolved_people_text']=text;item['record']['people']=[]
   warnings.append({'row':item['source_rows'][0],'message':'Plusieurs lignes pour un dossier : personnes à rapprocher, effectifs déclarés conservés une seule fois'})
  try:
   record=Register.model_validate(item['record'])
   for issue in quality(record,item['core']):
    warnings.append({'row':item['source_rows'][0],'message':'Avant soumission : '+issue['message']})
  except ValueError as e:errors.append({'row':item['source_rows'][0],'errors':[str(e)]})
 if not rows:warnings.append({'row':None,'message':'Modèle reconnu mais aucune ligne source renseignée ; aucun dossier fictif ne sera créé'})
 if date_role=='rp' and title=='Feuil1':warnings.append({'row':None,'message':'Choix explicite : Date de Feuil1 interprétée comme Date RP ; vérifier cette convention avant confirmation'})
 return {'date_role':date_role,'source_sheet':title,'sheets':sheets,'coverage':coverage,'total_rows':len(rows),'dossier_count':len(items),'items':items,'source_rows':source_rows,'errors':errors,'warnings':warnings,'file_sha256':hashlib.sha256(content).hexdigest()}

def preview_package(content):
 """Typed synthetic package: preserve relationships, ignore requested workflow status."""
 from .main import CaseIn
 lines=[line for line in content.decode('utf-8-sig').splitlines() if line.strip()]
 if len(lines)>10000:raise ValueError('Maximum 10 000 dossiers')
 items=[];sources=[];errors=[];seen=set();coverage={}
 for n,line in enumerate(lines,1):
  try:
   raw=json.loads(line)
   if raw.get('synthetic') is not True:raise ValueError('Paquet réservé aux données explicitement synthétiques')
   core=CaseIn.model_validate(raw['core']).model_dump(mode='json')
   if core['rp_year']!=date.fromisoformat(core['date_rp']).year:raise ValueError('Année RP différente de Date RP')
   key=(core['rp_number'].casefold(),core['rp_year'])
   if key in seen:raise ValueError('RP/année répété dans le paquet ; aucun rapprochement automatique')
   seen.add(key)
   record=validate_dates(Register.model_validate(raw['record']),date.fromisoformat(core['date_rp']))
   for prefix,obj in [('core',core),('record',record.model_dump(mode='json'))]:
    for k in obj:coverage[prefix+'.'+k]={'column':prefix+'.'+k,'destination':k,'handling':'mapped'}
   items.append({'core':core,'record':record.model_dump(mode='json'),'source_rows':[n]})
   sources.append({'sheet':'Registres JSONL','row':n,'cells':raw})
  except (ValueError,KeyError,TypeError,AttributeError) as e:errors.append({'row':n,'errors':[str(e)]})
 return {'date_role':None,'source_sheet':'Registres JSONL','sheets':[{'name':'Registres JSONL','rows':len(lines),'role':'source'}],'coverage':list(coverage.values()),'total_rows':len(lines),'dossier_count':len(items),'items':items,'source_rows':sources,'errors':errors,'warnings':[{'row':None,'message':'Paquet synthétique complet : relations conservées ; tout statut demandé dans le fichier est ignoré. Création en brouillon uniquement.'}],'file_sha256':hashlib.sha256(content).hexdigest()}

