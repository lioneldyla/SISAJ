"""Loss-aware reader and permanent provenance for the existing simplified importer."""
import base64
import csv
import hashlib
import io
import json
import zipfile
from pathlib import Path
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

PARSER_VERSION='legacy-trace-v1'
DDL='''
CREATE TABLE IF NOT EXISTS schema_migrations(version TEXT PRIMARY KEY, applied_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS import_sources(
 batch_id TEXT PRIMARY KEY REFERENCES import_batches(id),
 sha256 TEXT NOT NULL,
 original_base64 TEXT NOT NULL,
 parser_version TEXT NOT NULL,
 trace_json TEXT NOT NULL
);
'''

def checked_headers(values):
    headers=[str(v).strip() if v is not None else '' for v in values]
    used=[h.casefold() for h in headers if h]
    if not used or len(used)!=len(set(used)):
        raise HTTPException(400,'En-têtes absents ou dupliqués.')
    return headers

def read_source(filename,content):
    suffix=Path(filename).suffix.lower()
    records=[];trace=[]
    def add(headers,values,number,sheet):
        if not any(v is not None and str(v).strip() for v in values):return
        if len(records)>=10000:raise HTTPException(413,'Maximum 10 000 lignes par import.')
        if len(values)>len(headers) or any(not headers[i] and v is not None and str(v).strip() for i,v in enumerate(values)):
            raise HTTPException(400,'Valeur sans en-tête : aucune colonne ne peut être ignorée.')
        row={h:values[i] if i<len(values) else None for i,h in enumerate(headers) if h}
        records.append(row)
        trace.append({'sheet':sheet,'row':number,'cells':[
            {'column':get_column_letter(i+1),'header':h,'value':values[i] if i<len(values) else None}
            for i,h in enumerate(headers) if h]})
    if suffix=='.csv':
        try:
            reader=csv.reader(io.StringIO(content.decode('utf-8-sig'),newline=''),strict=True)
            headers=checked_headers(next(reader,[]))
            if len(headers)>256:raise HTTPException(413,'Maximum 256 colonnes.')
            # line_num includes blank lines and embedded newlines; store start line.
            while True:
                number=reader.line_num+1
                try:values=next(reader)
                except StopIteration:break
                add(headers,values,number,'CSV')
        except (UnicodeDecodeError,csv.Error) as exc:
            raise HTTPException(400,'CSV illisible ; UTF-8 et structure régulière requis.') from exc
        sheet=None;header_row=1
    elif suffix=='.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                members=z.infolist()
                if len(members)>2000 or sum(x.file_size for x in members)>60*1024*1024:
                    raise HTTPException(413,'Classeur décompressé trop volumineux.')
            wb=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
        except HTTPException:raise
        except Exception as exc:
            raise HTTPException(400,'XLSX illisible.') from exc
        try:
            ws=wb['AUDIENCES PENALES'] if 'AUDIENCES PENALES' in wb.sheetnames else wb.active
            sheet=ws.title;header_row=3 if sheet=='AUDIENCES PENALES' else 1
            if (ws.max_column or 0)>256 or (ws.max_row or 0)>10003:
                raise HTTPException(413,'Maximum 256 colonnes et 10 000 lignes.')
            header_cells=next(ws.iter_rows(min_row=header_row,max_row=header_row),[])
            headers=checked_headers([c.value for c in header_cells])
            for number,cells in enumerate(ws.iter_rows(min_row=header_row+1),header_row+1):
                if any(c.data_type=='f' for c in cells):
                    raise HTTPException(422,f'Formule à la ligne {number} : utiliser des valeurs source contrôlées.')
                add(headers,[c.value for c in cells],number,sheet)
        except HTTPException:raise
        except Exception as exc:
            raise HTTPException(400,'Contenu XLSX illisible.') from exc
        finally:wb.close()
    else:raise HTTPException(415,'Formats acceptés : CSV et XLSX.')
    return records,sheet,{'header_row':header_row,'rows':trace}

def preview_source(content,trace):
    return {'sha256':hashlib.sha256(content).hexdigest(),'original_base64':base64.b64encode(content).decode(),
            'parser_version':PARSER_VERSION,'trace':trace}

def persist(c,batch_id,payload):
    source=payload.get('provenance')
    if not source:
        raise HTTPException(409,'Ancien aperçu sans traçabilité : refaire la prévisualisation.')
    content=base64.b64decode(source['original_base64'],validate=True)
    if hashlib.sha256(content).hexdigest()!=source['sha256']:
        raise HTTPException(409,'Empreinte de la source incorrecte ; import annulé.')
    trace=source['trace']
    cases=c.execute('SELECT id FROM cases WHERE source_batch=? ORDER BY id',(batch_id,)).fetchall()
    if len(cases)!=len(trace['rows']):
        raise HTTPException(409,'Rapprochement des lignes source impossible ; import annulé.')
    trace={**trace,'rows':[{**row,'case_id':case['id']} for row,case in zip(trace['rows'],cases)]}
    c.execute('INSERT INTO import_sources VALUES(?,?,?,?,?)',(
        batch_id,source['sha256'],source['original_base64'],source['parser_version'],
        json.dumps(trace,ensure_ascii=False,default=str)))
