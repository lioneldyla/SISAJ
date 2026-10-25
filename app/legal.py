"""Additive v0.4 legal records. Method definitions are technical drafts, never official approval."""
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Literal
import json
import re

from fastapi import Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, ConfigDict, Field, model_validator

METHOD = 'sisa-legal-v0.4-draft'
SCHEMA = '''
CREATE TABLE IF NOT EXISTS jurisdictions(code TEXT PRIMARY KEY, name TEXT NOT NULL, created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS case_profiles(case_id INTEGER PRIMARY KEY REFERENCES cases(id), jurisdiction_code TEXT REFERENCES jurisdictions(code));
CREATE TABLE IF NOT EXISTS case_people(case_id INTEGER NOT NULL REFERENCES cases(id), reference TEXT NOT NULL, kind TEXT NOT NULL CHECK(kind IN ('physical','legal')), role TEXT NOT NULL CHECK(role IN ('defendant','victim','other')), display_name TEXT NOT NULL, PRIMARY KEY(case_id,reference));
CREATE TABLE IF NOT EXISTS case_offences(case_id INTEGER NOT NULL REFERENCES cases(id), reference TEXT NOT NULL, code TEXT NOT NULL, label TEXT NOT NULL, PRIMARY KEY(case_id,reference));
CREATE TABLE IF NOT EXISTS case_decisions(case_id INTEGER NOT NULL REFERENCES cases(id), reference TEXT NOT NULL, decision_date TEXT NOT NULL, on_merits INTEGER NOT NULL CHECK(on_merits IN (0,1)), PRIMARY KEY(case_id,reference));
CREATE TABLE IF NOT EXISTS decision_outcomes(case_id INTEGER NOT NULL, decision_reference TEXT NOT NULL, person_reference TEXT NOT NULL, offence_reference TEXT NOT NULL, result TEXT NOT NULL CHECK(result IN ('conviction','acquittal','other','unknown')), PRIMARY KEY(case_id,decision_reference,person_reference,offence_reference), FOREIGN KEY(case_id,decision_reference) REFERENCES case_decisions(case_id,reference), FOREIGN KEY(case_id,person_reference) REFERENCES case_people(case_id,reference), FOREIGN KEY(case_id,offence_reference) REFERENCES case_offences(case_id,reference));
CREATE TABLE IF NOT EXISTS case_measures(case_id INTEGER NOT NULL REFERENCES cases(id), reference TEXT NOT NULL, decision_reference TEXT, kind TEXT NOT NULL CHECK(kind IN ('fine','seizure','confiscation')), measure_date TEXT NOT NULL, amount TEXT, currency TEXT, PRIMARY KEY(case_id,reference), FOREIGN KEY(case_id,decision_reference) REFERENCES case_decisions(case_id,reference));
CREATE INDEX IF NOT EXISTS idx_decisions_date ON case_decisions(decision_date);
CREATE INDEX IF NOT EXISTS idx_measures_date ON case_measures(measure_date);
CREATE TABLE IF NOT EXISTS schema_migrations(version TEXT PRIMARY KEY, applied_at TEXT NOT NULL);
'''

class StrictModel(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True, extra='forbid')

class JurisdictionIn(StrictModel):
    code: str = Field(min_length=1,max_length=40,pattern=r'^[A-Za-z0-9_-]+$')
    name: str = Field(min_length=1,max_length=200)

class PersonIn(StrictModel):
    reference: str = Field(min_length=1,max_length=80)
    kind: Literal['physical','legal']
    role: Literal['defendant','victim','other']
    display_name: str = Field(min_length=1,max_length=200)

class OffenceIn(StrictModel):
    reference: str = Field(min_length=1,max_length=80)
    code: str = Field(min_length=1,max_length=80)
    label: str = Field(min_length=1,max_length=250)

class OutcomeIn(StrictModel):
    person_reference: str = Field(min_length=1,max_length=80)
    offence_reference: str = Field(min_length=1,max_length=80)
    result: Literal['conviction','acquittal','other','unknown']

class DecisionIn(StrictModel):
    reference: str = Field(min_length=1,max_length=80)
    decision_date: date
    on_merits: bool
    outcomes: list[OutcomeIn] = Field(default_factory=list,max_length=1000)

class MeasureIn(StrictModel):
    reference: str = Field(min_length=1,max_length=80)
    decision_reference: str | None = Field(default=None,min_length=1,max_length=80)
    kind: Literal['fine','seizure','confiscation']
    measure_date: date
    amount: str | None = Field(default=None,max_length=40)
    currency: str | None = Field(default=None,pattern=r'^[A-Z]{3}$')

    @model_validator(mode='after')
    def money(self):
        if (self.amount is None) != (self.currency is None):
            raise ValueError('Montant et devise sont requis ensemble ; inconnu = les deux absents.')
        if self.amount is not None:
            # Bounded fixed decimal, stored as exact canonical text, never float.
            if not re.fullmatch(r'\d{1,18}(\.\d{1,2})?', self.amount):
                raise ValueError('Montant positif ou nul : 18 chiffres et 2 décimales maximum.')
            self.amount = format(Decimal(self.amount).quantize(Decimal('0.01')), 'f')
        return self

class LegalIn(StrictModel):
    expected_version: int = Field(ge=1)
    reason: str = Field(min_length=5,max_length=1000)
    jurisdiction_code: str | None = Field(default=None,min_length=1,max_length=40)
    people: list[PersonIn] = Field(default_factory=list,max_length=1000)
    offences: list[OffenceIn] = Field(default_factory=list,max_length=1000)
    decisions: list[DecisionIn] = Field(default_factory=list,max_length=1000)
    measures: list[MeasureIn] = Field(default_factory=list,max_length=1000)

    @model_validator(mode='after')
    def links(self):
        for name in ('people','offences','decisions','measures'):
            items=getattr(self,name)
            if len({x.reference for x in items}) != len(items):
                raise ValueError('Référence dupliquée : '+name)
        people={p.reference:p for p in self.people}; offences={o.reference for o in self.offences}
        decisions={d.reference:d for d in self.decisions}
        for d in self.decisions:
            seen=set()
            for o in d.outcomes:
                if o.person_reference not in people or o.offence_reference not in offences:
                    raise ValueError('Résultat sans personne ou infraction du dossier.')
                if people[o.person_reference].role != 'defendant':
                    raise ValueError('Un résultat pénal doit concerner une personne poursuivie.')
                key=(o.person_reference,o.offence_reference)
                if key in seen: raise ValueError('Résultat dupliqué pour personne/infraction/décision.')
                seen.add(key)
        for m in self.measures:
            if m.decision_reference is not None and m.decision_reference not in decisions:
                raise ValueError('Mesure sans décision du dossier.')
            if m.decision_reference and m.measure_date < decisions[m.decision_reference].decision_date:
                raise ValueError('Mesure antérieure à sa décision de rattachement.')
        return self


def legal_snapshot(c,case_id):
    profile=c.execute('SELECT jurisdiction_code FROM case_profiles WHERE case_id=?',(case_id,)).fetchone()
    def rows(table):
        return [dict(r) for r in c.execute('SELECT * FROM '+table+' WHERE case_id=? ORDER BY reference',(case_id,)).fetchall()]
    people=rows('case_people'); offences=rows('case_offences'); decisions=rows('case_decisions'); measures=rows('case_measures')
    for values in (people,offences,decisions,measures):
        for r in values: r.pop('case_id')
    for d in decisions:
        d['on_merits']=bool(d['on_merits'])
        d['outcomes']=[dict(r) for r in c.execute('SELECT person_reference,offence_reference,result FROM decision_outcomes WHERE case_id=? AND decision_reference=? ORDER BY person_reference,offence_reference',(case_id,d['reference'])).fetchall()]
    return {'jurisdiction_code':profile['jurisdiction_code'] if profile else None,'people':people,'offences':offences,'decisions':decisions,'measures':measures}


def calculate(c,start=None,end=None,jurisdiction=None):
    if start and end and start>end: raise HTTPException(422,'Période invalide.')
    args=[]; wh="c.status='validated'"
    if jurisdiction:
        wh+=' AND EXISTS (SELECT 1 FROM case_profiles p WHERE p.case_id=c.id AND p.jurisdiction_code=?)';args.append(jurisdiction)
    def period(column):
        clause=''; values=[]
        if start: clause+=' AND '+column+'>=?';values.append(start.isoformat())
        if end: clause+=' AND '+column+'<=?';values.append(end.isoformat())
        return clause,values
    pc,pa=period('c.date_rp'); pd,da=period('d.decision_date'); pm,ma=period('m.measure_date')
    received=c.execute('SELECT COUNT(*) n FROM cases c WHERE '+wh+pc,args+pa).fetchone()['n']
    # One decision counted once regardless of people/offences/outcomes.
    decisions=c.execute('SELECT d.*,c.rp_number FROM case_decisions d JOIN cases c ON c.id=d.case_id WHERE '+wh+pd+' ORDER BY d.case_id,d.reference',args+da).fetchall()
    convicted=set();acquitted=set();mixed=set();judged=set();unknown=0
    for d in decisions:
        if not d['on_merits']: continue
        key=(d['case_id'],d['reference']);judged.add(d['case_id'])
        results={r['result'] for r in c.execute('SELECT result FROM decision_outcomes WHERE case_id=? AND decision_reference=?',(d['case_id'],d['reference'])).fetchall()}
        if 'conviction' in results: convicted.add(key)
        if 'acquittal' in results: acquitted.add(key)
        if {'conviction','acquittal'} <= results: mixed.add(key)
        if not results or 'unknown' in results: unknown+=1
    amounts={}; missing=0
    measures=c.execute('SELECT m.* FROM case_measures m JOIN cases c ON c.id=m.case_id WHERE '+wh+pm+' ORDER BY m.case_id,m.reference',args+ma).fetchall()
    for m in measures:
        if m['amount'] is None: missing+=1;continue
        key=(m['kind'],m['currency']);amounts[key]=amounts.get(key,Decimal('0'))+Decimal(m['amount'])
    return {'method_version':METHOD,'officially_approved':False,'period':{'start':str(start) if start else None,'end':str(end) if end else None},'jurisdiction_code':jurisdiction,
      'population':'dossiers validés uniquement','received_cases':received,'judged_cases':len(judged),'decisions':len(decisions),
      'merits_decisions':sum(bool(d['on_merits']) for d in decisions),'conviction_decisions':len(convicted),'acquittal_decisions':len(acquitted),'mixed_decisions':len(mixed),'unknown_outcome_decisions':unknown,
      'measures':len(measures),'unknown_amount_measures':missing,'amounts':[{'kind':k[0],'currency':k[1],'amount':format(v,'f')} for k,v in sorted(amounts.items())],
      'date_basis':{'received_cases':'date_rp','decisions':'decision_date','judged_cases':'decision_date','measures':'measure_date'},
      'warning':'Méthode technique provisoire. Condamnations et relaxes peuvent se recouvrir ; aucune somme entre devises. Événements historiques et effectifs déclarés non convertis automatiquement.'}


def install_legal(app,db,locked_case,audit,utcnow,user_from_token,require_roles,safe_cell):
    with db() as c:
        c.executescript(SCHEMA)
        c.execute('INSERT OR IGNORE INTO schema_migrations(version,applied_at) VALUES (?,?)',('004-legal-additive',utcnow()))

    @app.get('/jurisdictions')
    def jurisdictions(u=Depends(user_from_token)):
        with db() as c: return {'items':[dict(r) for r in c.execute('SELECT code,name FROM jurisdictions ORDER BY code').fetchall()]}

    @app.post('/jurisdictions')
    def add_jurisdiction(x:JurisdictionIn,u=Depends(require_roles('admin','supervisor'))):
        with db(write=True) as c:
            c.execute('INSERT INTO jurisdictions(code,name,created_at) VALUES (?,?,?)',(x.code,x.name,utcnow()))
            audit(c,u['id'],'create','jurisdiction',x.code,x.model_dump())
        return x.model_dump()

    @app.get('/cases/{case_id}/legal')
    def get_legal(case_id:int,u=Depends(user_from_token)):
        with db() as c:
            row=c.execute('SELECT version FROM cases WHERE id=?',(case_id,)).fetchone()
            if not row: raise HTTPException(404,'Dossier introuvable.')
            return {'case_id':case_id,'version':row['version'],**legal_snapshot(c,case_id)}

    @app.put('/cases/{case_id}/legal')
    def save_legal(case_id:int,x:LegalIn,u=Depends(require_roles('admin','collector','supervisor'))):
        with db(write=True) as c:
            row=locked_case(c,case_id,x.expected_version)
            if row['status'] not in ('draft','returned'):
                raise HTTPException(409,'Remettre le dossier en brouillon par une correction motivée avant de modifier son contenu juridique.')
            if x.jurisdiction_code and not c.execute('SELECT code FROM jurisdictions WHERE code=?',(x.jurisdiction_code,)).fetchone():
                raise HTTPException(422,'Juridiction inconnue.')
            rp=date.fromisoformat(row['date_rp'])
            if any(d.decision_date<rp for d in x.decisions) or any(m.measure_date<rp for m in x.measures):
                raise HTTPException(422,'Décision ou mesure antérieure à la date RP.')
            before=legal_snapshot(c,case_id)
            for table in ('decision_outcomes','case_measures','case_decisions','case_offences','case_people','case_profiles'):
                c.execute('DELETE FROM '+table+' WHERE case_id=?',(case_id,))
            c.execute('INSERT INTO case_profiles(case_id,jurisdiction_code) VALUES (?,?)',(case_id,x.jurisdiction_code))
            for p in x.people:
                c.execute('INSERT INTO case_people(case_id,reference,kind,role,display_name) VALUES (?,?,?,?,?)',(case_id,p.reference,p.kind,p.role,p.display_name))
            for o in x.offences:
                c.execute('INSERT INTO case_offences(case_id,reference,code,label) VALUES (?,?,?,?)',(case_id,o.reference,o.code,o.label))
            for d in x.decisions:
                c.execute('INSERT INTO case_decisions(case_id,reference,decision_date,on_merits) VALUES (?,?,?,?)',(case_id,d.reference,d.decision_date.isoformat(),int(d.on_merits)))
                for o in d.outcomes:
                    c.execute('INSERT INTO decision_outcomes(case_id,decision_reference,person_reference,offence_reference,result) VALUES (?,?,?,?,?)',(case_id,d.reference,o.person_reference,o.offence_reference,o.result))
            for m in x.measures:
                c.execute('INSERT INTO case_measures(case_id,reference,decision_reference,kind,measure_date,amount,currency) VALUES (?,?,?,?,?,?,?)',(case_id,m.reference,m.decision_reference,m.kind,m.measure_date.isoformat(),m.amount,m.currency))
            c.execute('UPDATE cases SET version=version+1,updated_by=?,updated_at=? WHERE id=?',(u['id'],utcnow(),case_id))
            audit(c,u['id'],'update','legal_record',case_id,{'reason':x.reason,'before':before,'after':legal_snapshot(c,case_id),'version':row['version']+1})
        return {'case_id':case_id,'version':row['version']+1}

    @app.get('/indicators/catalogue')
    def catalogue(u=Depends(user_from_token)):
        return {'method_version':METHOD,'officially_approved':False,'indicators':[
            {'code':'received_cases','unit':'dossier distinct','date_basis':'date_rp','rule':'Dossiers validés dont la date RP appartient à la période.'},
            {'code':'judged_cases','unit':'dossier distinct','date_basis':'decision_date','rule':'Au moins une décision au fond dans la période ; chaque dossier compté une fois.'},
            {'code':'merits_decisions','unit':'décision distincte dans son dossier','date_basis':'decision_date','rule':'Décisions marquées au fond dans la période.'},
            {'code':'conviction_decisions','unit':'décision distincte','date_basis':'decision_date','rule':'Décision au fond avec au moins un résultat condamnation.'},
            {'code':'acquittal_decisions','unit':'décision distincte','date_basis':'decision_date','rule':'Décision au fond avec au moins un résultat relaxe.'},
            {'code':'mixed_decisions','unit':'décision distincte','date_basis':'decision_date','rule':'Même décision au fond contenant condamnation et relaxe ; incluse dans les deux indicateurs.'},
            {'code':'amounts','unit':'montant par nature et devise','date_basis':'measure_date','rule':'Somme décimale exacte des montants connus ; inconnu distinct de zéro ; devises jamais additionnées.'}]}

    @app.get('/stats/legal')
    def legal_stats(start:date|None=None,end:date|None=None,jurisdiction:str|None=None,u=Depends(user_from_token)):
        with db() as c: return calculate(c,start,end,jurisdiction)

    @app.get('/reports/legal.xlsx')
    def legal_xlsx(start:date|None=None,end:date|None=None,jurisdiction:str|None=None,u=Depends(require_roles('admin','supervisor'))):
        from openpyxl import Workbook
        with db() as c:
            result=calculate(c,start,end,jurisdiction)
            audit(c,u['id'],'export','legal_report',detail={'format':'xlsx','method':METHOD,'start':start,'end':end,'jurisdiction':jurisdiction})
        wb=Workbook();ws=wb.active;ws.title='Indicateurs'
        ws.append(['Indicateur','Valeur'])
        for key in ('received_cases','judged_cases','decisions','merits_decisions','conviction_decisions','acquittal_decisions','mixed_decisions','unknown_outcome_decisions','measures','unknown_amount_measures'):ws.append([key,result[key]])
        amounts=wb.create_sheet('Montants');amounts.append(['Nature','Devise','Montant exact (texte)'])
        for r in result['amounts']:amounts.append([r['kind'],r['currency'],r['amount']])
        meta=wb.create_sheet('Méthode')
        for key,val in [('Méthode',METHOD),('Approbation officielle','Non'),('Début',str(start or 'Toutes dates')),('Fin',str(end or 'Toutes dates')),('Juridiction',jurisdiction or 'Toutes'),('Population',result['population']),('Avertissement',result['warning']),('Dates',json.dumps(result['date_basis'],ensure_ascii=False)),('Auteur',u['username']),('UTC',utcnow())]:meta.append([key,safe_cell(val)])
        buf=BytesIO();wb.save(buf);buf.seek(0)
        return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="sisa_indicateurs.xlsx"'})

    @app.get('/reports/legal.pdf')
    def legal_pdf(start:date|None=None,end:date|None=None,jurisdiction:str|None=None,u=Depends(require_roles('admin','supervisor'))):
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import simpleSplit
        with db() as c:
            result=calculate(c,start,end,jurisdiction)
            audit(c,u['id'],'export','legal_report',detail={'format':'pdf','method':METHOD,'start':start,'end':end,'jurisdiction':jurisdiction})
        buf=BytesIO();pdf=canvas.Canvas(buf);pdf.setTitle('SISA - Indicateurs juridiques provisoires');y=790
        lines=['SISA - Rapport juridique de démonstration',METHOD,'Non approuvé institutionnellement',f"Période : {start or 'Toutes dates'} / {end or 'Toutes dates'}",f"Juridiction : {jurisdiction or 'Toutes'}",'Population : dossiers validés uniquement',f"Auteur : {u['username']} | UTC : {utcnow()}"]
        lines += [f'{k} : {result[k]}' for k in ('received_cases','judged_cases','merits_decisions','conviction_decisions','acquittal_decisions','mixed_decisions','unknown_outcome_decisions','unknown_amount_measures')]
        lines += [f"{r['kind']} : {r['amount']} {r['currency']}" for r in result['amounts']]
        lines += [json.dumps(result['date_basis'],ensure_ascii=False),result['warning']]
        for line in lines:
            for part in simpleSplit(line,'Helvetica',10,490):
                if y<55:pdf.showPage();y=790
                pdf.setFont('Helvetica',10);pdf.drawString(45,y,part);y-=18
        pdf.save();buf.seek(0)
        return StreamingResponse(buf,media_type='application/pdf',headers={'Content-Disposition':'attachment; filename="sisa_indicateurs.pdf"'})
