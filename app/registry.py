"""Complete PPEF register, source-preserving transactional import and six operational statements."""
import base64,hashlib,io,json,secrets,time,uuid,zipfile
from collections import Counter,defaultdict
from datetime import date,datetime
from decimal import Decimal
from statistics import mean,median
from fastapi import Depends,File,Form,HTTPException,UploadFile
from fastapi.responses import StreamingResponse
from .registry_models import Register,SaveRegister,CATALOGUES,quality,validate_dates
from . import ppef_import

SCHEMA='''
CREATE TABLE IF NOT EXISTS registry_records(case_id INTEGER PRIMARY KEY REFERENCES cases(id), payload TEXT NOT NULL, updated_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS registry_previews(id TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), created_at INTEGER NOT NULL, payload TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS registry_batches(id TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), created_at TEXT NOT NULL, filename TEXT NOT NULL, sha256 TEXT NOT NULL, original_base64 TEXT NOT NULL, manifest TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS registry_source_rows(id TEXT PRIMARY KEY, batch_id TEXT NOT NULL REFERENCES registry_batches(id), case_id INTEGER NOT NULL REFERENCES cases(id), sheet TEXT NOT NULL, row_number INTEGER NOT NULL, payload TEXT NOT NULL);
CREATE INDEX IF NOT EXISTS idx_registry_source_case ON registry_source_rows(case_id);
'''
METHOD='ppef-register-v0.5-draft'

def read_record(c,case):
 r=c.execute('SELECT payload FROM registry_records WHERE case_id=?',(case['id'],)).fetchone()
 if r:return Register.model_validate_json(r['payload'])
 from .legal import legal_snapshot
 old=legal_snapshot(c,case['id'])
 return Register.model_validate({'jurisdiction_code':old['jurisdiction_code'],'ri_state':'assigned' if case['ri_number'] else 'not_recorded','transnational_state':'yes' if case['transnational'] else 'no',
 'people':old['people'],
 'offences':[{'reference':o['reference'],'parquet_label':o['label']} for o in old['offences']],
 'decisions':[{'reference':d['reference'],'number':d['reference'],'decision_date':d['decision_date'],'nature':'merits' if d['on_merits'] else 'not_merits','result':'not_recorded','outcomes':d['outcomes']} for d in old['decisions']],
 'measures':[{'reference':m['reference'],'kind':m['kind'],'decision_reference':m['decision_reference'],'measure_date':m['measure_date'],'amount':{'state':'known' if m['amount'] is not None else 'not_recorded','value':m['amount'],'currency':m['currency'] or 'XOF'}} for m in old['measures']]})

def write_record(c,cid,r,now,user):
 encoded=r.model_dump_json()
 if c.execute('SELECT case_id FROM registry_records WHERE case_id=?',(cid,)).fetchone():c.execute('UPDATE registry_records SET payload=?,updated_at=? WHERE case_id=?',(encoded,now,cid))
 else:c.execute('INSERT INTO registry_records(case_id,payload,updated_at) VALUES (?,?,?)',(cid,encoded,now))
 if r.count_basis=='enumerated':
  pp=sum(p.kind=='physical' and p.role=='defendant' for p in r.people);pm=sum(p.kind=='legal' and p.role=='defendant' for p in r.people)
  c.execute('UPDATE cases SET person_physical_count=?,person_legal_count=? WHERE id=?',(pp,pm,cid))
 # Unknown transnationality stays explicit in the typed record. Legacy storage is a projection only.
 c.execute('UPDATE cases SET transnational=?,version=version+1,updated_at=?,updated_by=? WHERE id=?',(int(r.transnational_state=='yes'),now,user,cid))

def ready_for_submission(c,cid):
 row=c.execute('SELECT * FROM cases WHERE id=?',(cid,)).fetchone()
 record=c.execute('SELECT payload FROM registry_records WHERE case_id=?',(cid,)).fetchone()
 if not record:return
 issues=quality(Register.model_validate_json(record['payload']),dict(row))
 if issues:raise HTTPException(422,'Compléter le registre avant soumission : '+' ; '.join(x['message'] for x in issues))

def operational(c,start=None,end=None,as_of=None,cabinet=None,category=None,provenance=None,waiting_days=90,jurisdiction=None):
 if start and end and start>end:raise HTTPException(422,'Période invalide')
 as_of=as_of or end or date.today()
 if start and start>as_of:raise HTTPException(422,'Date d’observation antérieure au début de période')
 def in_period(d):return bool(d and (not start or str(d)>=str(start)) and (not end or str(d)<=str(end)) and str(d)<=str(as_of))
 cases=[dict(x) for x in c.execute("SELECT * FROM cases WHERE status='validated' ORDER BY id").fetchall()]
 records={row['case_id']:Register.model_validate_json(row['payload']) for row in c.execute('SELECT case_id,payload FROM registry_records').fetchall()}
 cats=Counter();units=Counter();cabinets=Counter();entries=Counter();exits=Counter();groups=Counter();counts=Counter();amounts=defaultdict(Decimal);missing=Counter();basis=Counter();issues=[];delays=defaultdict(list);stocks=[];persons=Counter();declared=Counter();sex=Counter();nationality=Counter();sentences=[];fine_values=defaultdict(list)
 selected=[]
 for case in cases:
  r=records.get(case['id']) or read_record(c,case);p=r.procedure
  if jurisdiction and r.jurisdiction_code!=jurisdiction:continue
  if cabinet and p.cabinet!=cabinet:continue
  if category and case['category']!=category:continue
  if provenance and case['provenance']!=provenance:continue
  received=date.fromisoformat(case['date_rp']);closed=p.closure_date
  if received>as_of:continue
  selected.append((case,r))
  if not closed or closed>as_of:stocks.append({'id':case['id'],'rp_number':case['rp_number'],'cabinet':p.cabinet,'age_days':(as_of-received).days})
  if in_period(received):
   counts['received']+=1;entries[str(received)[:7]]+=1;cats[case['category']]+=1;units[case['provenance']]+=1;cabinets[p.cabinet or 'Non affecté']+=1
   counts['transnational_'+r.transnational_state]+=1;basis[r.count_basis]+=1
   for key,field in [('physical','person_physical_count'),('legal','person_legal_count')]:
    if r.count_basis=='enumerated':persons[key]+=sum(x.kind==key and x.role=='defendant' for x in r.people)
    elif case[field] is not None:declared[key]+=case[field]
    else:missing['effectif_'+key]+=1
   for person in r.people:
    if person.role=='defendant' and person.public_agent=='yes':persons['public_agents']+=1
   for field in ('convicted_total','convicted_male','convicted_female','victim_total'):
    val=getattr(r.declarations,field)
    if val is not None:declared[field]+=val
   for group in {o.offence_group for o in r.offences}:groups[group]+=1
   for o in r.offences:
    if o.amount.state=='known':amounts[('offence',o.amount.currency,o.amount_basis)]+=Decimal(o.amount.value)
    else:missing['offence_amount']+=1
  if in_period(closed):counts['closed']+=1;exits[str(closed)[:7]]+=1;counts['closure_'+p.closure_kind]+=1;delays['receipt_to_closure'].append((closed-received).days)
  if p.first_hearing_held and in_period(p.first_hearing_held):delays['receipt_to_hearing'].append((p.first_hearing_held-received).days)
  judged=False;convicted=False;acquitted=False
  for d in r.decisions:
   if not in_period(d.decision_date):
    if d.decision_date is None:missing['decision_date']+=1
    continue
   counts['decisions']+=1
   is_merits=d.nature in ('merits','mixed','partial_merits')
   if is_merits:
    judged=True;counts['merits_decisions']+=1
    resultset={x.result for x in d.outcomes}
    has_conv='conviction' in resultset or d.result in ('conviction','mixed');has_acq='acquittal' in resultset or d.result in ('acquittal','mixed')
    convicted|=has_conv;acquitted|=has_acq
    counts['conviction_decisions']+=has_conv;counts['acquittal_decisions']+=has_acq;counts['mixed_decisions']+=has_conv and has_acq
    if (not d.outcomes and d.result in ('not_recorded','undetermined','pending','not_rendered')) or any(o.result=='unknown' for o in d.outcomes):counts['unknown_outcome_decisions']+=1
    convicted_people={x.person_reference for x in d.outcomes if x.result=='conviction'}
    for person in r.people:
     if person.reference in convicted_people:sex[person.sex]+=1;nationality[person.nationality or 'Non renseignée']+=1
    for s in r.sentences:
     if s.decision_reference==d.reference and s.duration_months is not None:sentences.append(s.duration_months)
    delays['receipt_to_decision'].append((d.decision_date-received).days)
   else:counts['non_merits_or_unknown_decisions']+=1
  counts['judged_cases']+=judged;counts['convicted_cases']+=convicted;counts['acquitted_cases']+=acquitted
  seizure=False;confiscation=False
  for m in r.measures:
   if not in_period(m.measure_date):
    if m.measure_date is None:missing['measure_date']+=1
    continue
   counts[m.kind+'_measures']+=1;seizure|=m.kind=='seizure';confiscation|=m.kind=='confiscation'
   if m.amount.state=='known':
    amounts[(m.kind,m.amount.currency,m.valuation_basis)]+=Decimal(m.amount.value)
    if m.kind=='fine':fine_values[m.amount.currency].append(Decimal(m.amount.value))
   else:missing[m.kind+'_amount']+=1
  counts['seizure_cases']+=seizure;counts['confiscation_cases']+=confiscation
  if not closed or closed>as_of:
   if p.orientation in ('not_recorded','pending','undetermined'):issues.append({'case_id':case['id'],'code':'WAITING_ORIENTATION','message':'Orientation à examiner'})
   if not p.first_hearing_planned and not p.first_hearing_held:issues.append({'case_id':case['id'],'code':'HEARING_NOT_RECORDED','message':'Audience non renseignée ; vérifier si applicable'})
   if not any(d.decision_date for d in r.decisions):issues.append({'case_id':case['id'],'code':'DECISION_NOT_RECORDED','message':'Décision non renseignée ; vérifier le stade'})
   if (as_of-received).days>waiting_days:issues.append({'case_id':case['id'],'code':'AGE_THRESHOLD','message':f'Ancienneté supérieure au seuil indicatif de {waiting_days} jours'})
  issues.extend({'case_id':case['id'],**i} for i in quality(r,case))
 # A specified denominator, never a percentage of heterogeneous fields or an invented PASS.
 complete=0;expected=0
 for case,r in selected:
  fields=[bool(case['rp_number']),bool(case['date_rp']),case['provenance']!='Non renseigné',case['category']!='Qualification non déterminée',r.transnational_state in ('yes','no')]
  fields += [bool(d.number) and d.decision_date is not None for d in r.decisions if d.result not in ('not_rendered','pending')]
  fields += [m.measure_date is not None for m in r.measures]
  complete+=sum(fields);expected+=len(fields)
 def dist(v):return [{'label':k,'count':n} for k,n in sorted(v.items())]
 forecast={'available':False,'reason':'Au moins 24 mois calendaires d’historique validé sont requis. Aucun modèle de culpabilité ou de décision individuelle.'}
 all_months=Counter()
 for case,r in selected:
  if date.fromisoformat(case['date_rp'])<=as_of:all_months[case['date_rp'][:7]]+=1
 if all_months:
  first=min(all_months);y,m=map(int,first.split('-'));series=[]
  while f'{y:04}-{m:02}'<=str(as_of)[:7]:
   series.append((f'{y:04}-{m:02}',all_months[f'{y:04}-{m:02}']));m+=1
   if m==13:y+=1;m=1
  # Exclude the current incomplete month from fitting/backtest.
  series=[s for s in series if s[0]<str(as_of)[:7]]
  if len(series)>=24 and len(all_months)>=24:
   values=[v for _,v in series];predictions=[mean(values[i-6:i]) for i in range(len(values)-6,len(values))]
   mae=mean(abs(pred-values[len(values)-6+i]) for i,pred in enumerate(predictions));future=mean(values[-6:])
   forecast={'available':True,'method':'Moyenne mobile 6 mois ; base descriptive, non IA','history_months':len(series),'backtest_months':6,'mae':round(mae,2),'monthly_estimate_next_3_months':round(future,2),'warning':'Scénario indicatif sans intervalle calibré ; historique antérieur à la date d’observation, indépendamment de la période des flux ; mois incomplet exclu ; aucune estimation individuelle de culpabilité ou de durée.'}
 return {'method_version':METHOD,'officially_approved':False,'period':{'start':str(start) if start else None,'end':str(end) if end else None,'as_of':str(as_of)},'population':'Dossiers validés uniquement','counts':dict(counts),'stock':{'count':len(stocks),'items':stocks[:500],'total':len(stocks)},'by_category':dist(cats),'by_provenance':dist(units),'by_cabinet':dist(cabinets),'by_offence_group':dist(groups),'monthly':{'entries':dist(entries),'exits':dist(exits)},'persons':{'enumerated':dict(persons),'declared':dict(declared),'count_basis':dict(basis),'sex_convicted_outcomes':dict(sex),'nationalities_convicted_outcomes':dist(nationality),'unit':'Personne dans son dossier ; aucune déduplication inter-dossiers ni entre décisions'},'delays':{k:{'n':len(v),'mean_days':round(mean(v),2),'median_days':median(v)} for k,v in delays.items()},'financial':[{'kind':k[0],'currency':k[1],'basis':k[2],'amount':format(v,'f')} for k,v in sorted(amounts.items())],'sentence_months':{'min':min(sentences) if sentences else None,'max':max(sentences) if sentences else None},'fine_ranges':[{'currency':k,'min':str(min(v)),'max':str(max(v))} for k,v in fine_values.items()],'quality':{'issues':issues[:500],'total_issues':len(issues),'missing':dict(missing),'complete':complete,'expected':expected,'completeness_percent':round(100*complete/expected,2) if expected else None,'denominator':'5 champs de réception + référence/date de chaque décision rendue + date de chaque mesure'},'throughput_percent':round(100*counts['closed']/counts['received'],2) if counts['received'] else None,'forecast':forecast,'warnings':['Taux sorties/entrées distinct du taux de résolution d’une cohorte.','Dates : réception par Date RP ; clôtures par date de clôture ; décisions par date de décision ; mesures par date propre.','Sommes séparées par nature, devise et méthode de valorisation. Les inconnus ne sont pas des zéros.','Nomenclatures proposées, à valider par la juridiction.']}


def install(app,db,locked_case,audit,utcnow,user_from_token,require_roles,safe_cell):
 with db() as c:
  c.executescript(SCHEMA)
  c.execute('INSERT OR IGNORE INTO schema_migrations VALUES (?,?)',('005-complete-register',utcnow()))
 @app.get('/registry/catalogue')
 def catalogue(u=Depends(user_from_token)):
  return {'method_version':METHOD,'catalogues':CATALOGUES,'schema':Register.model_json_schema(),'approved':False,'import_fields':ppef_import.INPUT,'derived_columns':ppef_import.DERIVED}
 @app.get('/cases/{case_id}/register')
 def get(case_id:int,u=Depends(user_from_token)):
  with db(write=True) as c:
   case=dict(locked_case(c,case_id));r=read_record(c,case)
   audit(c,u['id'],'consult','register',case_id,{})
   sources=[dict(x) for x in c.execute('SELECT batch_id,sheet,row_number FROM registry_source_rows WHERE case_id=? ORDER BY sheet,row_number',(case_id,)).fetchall()]
  return {'case':case,'record':r.model_dump(mode='json'),'issues':quality(r,case),'sources':sources}
 @app.put('/cases/{case_id}/register')
 def save(case_id:int,x:SaveRegister,u=Depends(require_roles('admin','supervisor','collector'))):
  with db(write=True) as c:
   case=dict(locked_case(c,case_id,x.expected_version))
   if case['status'] not in ('draft','returned'):raise HTTPException(409,'Corriger le dossier avec un motif pour le remettre en brouillon avant modification')
   try:r=validate_dates(x.record,date.fromisoformat(case['date_rp']))
   except ValueError as e:raise HTTPException(422,str(e))
   if r.jurisdiction_code and not c.execute('SELECT code FROM jurisdictions WHERE code=?',(r.jurisdiction_code,)).fetchone():raise HTTPException(422,'Juridiction inconnue')
   before=read_record(c,case).model_dump(mode='json');write_record(c,case_id,r,utcnow(),u['id'])
   audit(c,u['id'],'update','register',case_id,{'reason':x.reason,'before':before,'after':r.model_dump(mode='json'),'version':case['version']+1})
  return {'id':case_id,'version':case['version']+1,'issues':quality(r,case)}
 @app.post('/registry/imports/preview')
 async def preview(file:UploadFile=File(...),sheet:str|None=Form(default=None),date_role:str|None=Form(default=None),u=Depends(require_roles('admin','supervisor','collector'))):
  content=await file.read(15*1024*1024+1)
  if len(content)>15*1024*1024:raise HTTPException(413,'Fichier limité à 15 Mo')
  if date_role not in (None,'','rp'):raise HTTPException(422,'Convention de date non reconnue')
  try:result=ppef_import.preview(file.filename or '',content,sheet or None,date_role or None)
  except Exception as e:
   if isinstance(e,(ValueError,UnicodeError,zipfile.BadZipFile)):raise HTTPException(422,str(e))
   raise
  with db(write=True) as c:
   for item in result['items']:
    if c.execute('SELECT id FROM cases WHERE LOWER(rp_number)=? AND rp_year=?',(item['core']['rp_number'].lower(),item['core']['rp_year'])).fetchone():result['errors'].append({'row':item['source_rows'][0],'errors':['RP/année déjà présent : aucun écrasement automatique']})
   pid=secrets.token_urlsafe(18)
   payload={'result':result,'filename':file.filename,'original_base64':base64.b64encode(content).decode()}
   c.execute('DELETE FROM registry_previews WHERE created_at<?',(int(time.time())-1800,))
   c.execute('INSERT INTO registry_previews VALUES (?,?,?,?)',(pid,u['id'],int(time.time()),json.dumps(payload,ensure_ascii=False,default=str)))
  return {'preview_id':pid,'source_sheet':result['source_sheet'],'sheets':result['sheets'],'coverage':result['coverage'],'total_rows':result['total_rows'],'valid_rows':len(result['items']),'error_count':len(result['errors']),'errors':result['errors'][:200],'warnings':result['warnings'][:200],'warning_count':len(result['warnings']),'can_confirm':bool(result['items']) and not result['errors'],'sample':[x['core'] for x in result['items'][:5]],'message':'Original complet archivé seulement à la confirmation ; calculs Excel non pris pour des faits.'}
 @app.post('/registry/imports/{preview_id}/confirm')
 def confirm(preview_id:str,u=Depends(require_roles('admin','supervisor','collector'))):
  if not __import__('re').fullmatch('[A-Za-z0-9_-]{24}',preview_id):raise HTTPException(404,'Aperçu introuvable')
  with db(write=True) as c:
   from .main import DATABASE_URL
   row=c.execute('SELECT * FROM registry_previews WHERE id=?'+(' FOR UPDATE' if DATABASE_URL else ''),(preview_id,)).fetchone()
   if not row:raise HTTPException(404,'Aperçu absent ou déjà confirmé')
   if row['user_id']!=u['id']:raise HTTPException(403,'Aperçu créé par un autre utilisateur')
   if row['created_at']<int(time.time())-1800:raise HTTPException(409,'Aperçu expiré ; contrôler le fichier à nouveau')
   payload=json.loads(row['payload']);r=payload['result']
   if r['errors'] or not r['items']:raise HTTPException(422,'Import bloqué : corriger les erreurs ou fournir des lignes source')
   if DATABASE_URL:c.execute('SELECT pg_advisory_xact_lock(731026)')
   for item in r['items']:
    if c.execute('SELECT id FROM cases WHERE LOWER(rp_number)=? AND rp_year=?',(item['core']['rp_number'].lower(),item['core']['rp_year'])).fetchone():raise HTTPException(409,'Un dossier du lot existe désormais ; aucune ligne importée, refaire le contrôle.')
   bid=str(uuid.uuid4());now=utcnow()
   c.execute('INSERT INTO registry_batches VALUES (?,?,?,?,?,?,?)',(bid,u['id'],now,payload['filename'],r['file_sha256'],payload['original_base64'],json.dumps({k:r[k] for k in ('sheets','coverage','warnings')},ensure_ascii=False)))
   for item in r['items']:
    d=item['core'];cur=c.execute('INSERT INTO cases(rp_number,rp_year,date_rp,provenance,category,ri_number,transnational,person_physical_count,person_legal_count,source_batch,created_by,updated_by,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(d['rp_number'],d['rp_year'],d['date_rp'],d['provenance'],d['category'],d['ri_number'],int(d['transnational']),d['person_physical_count'],d['person_legal_count'],bid,u['id'],u['id'],now,now));cid=cur.lastrowid
    record=Register.model_validate(item['record']);write_record(c,cid,record,now,u['id'])
    for source in r['source_rows']:
     if source['row'] in item['source_rows']:c.execute('INSERT INTO registry_source_rows VALUES (?,?,?,?,?,?)',(str(uuid.uuid4()),bid,cid,source['sheet'],source['row'],json.dumps(source['cells'],ensure_ascii=False)))
   c.execute('INSERT INTO import_batches VALUES (?,?,?,?,?,?,?)',(bid,payload['filename'],r['total_rows'],len(r['items']),u['id'],now,r['source_sheet']))
   audit(c,u['id'],'import','registry_batch',bid,{'rows':r['total_rows'],'dossiers':len(r['items']),'sha256':r['file_sha256']})
   c.execute('DELETE FROM registry_previews WHERE id=?',(preview_id,))
  return {'status':'imported','batch_id':bid,'imported_rows':len(r['items']),'source_rows':r['total_rows']}
 @app.get('/registry/imports/{batch_id}/source')
 def original(batch_id:str,u=Depends(require_roles('admin','supervisor'))):
  with db(write=True) as c:
   r=c.execute('SELECT * FROM registry_batches WHERE id=?',(batch_id,)).fetchone()
   if not r:raise HTTPException(404,'Lot introuvable')
   content=base64.b64decode(r['original_base64']);assert hashlib.sha256(content).hexdigest()==r['sha256']
   audit(c,u['id'],'export','registry_source',batch_id,{})
  suffix='.xlsx' if r['filename'].lower().endswith('.xlsx') else '.csv'
  return StreamingResponse(io.BytesIO(content),media_type='application/octet-stream',headers={'Content-Disposition':'attachment; filename="source_ppef'+suffix+'"'})
 @app.get('/registry/statements')
 def statements(start:date|None=None,end:date|None=None,as_of:date|None=None,cabinet:str|None=None,category:str|None=None,provenance:str|None=None,waiting_days:int=90,u=Depends(user_from_token)):
  if not 1<=waiting_days<=3650:raise HTTPException(422,'Seuil indicatif : 1 à 3650 jours')
  with db() as c:return operational(c,start,end,as_of,cabinet,category,provenance,waiting_days)
 @app.get('/registry/export.json')
 def export(u=Depends(require_roles('admin','supervisor'))):
  with db(write=True) as c:
   cases=[dict(r) for r in c.execute('SELECT * FROM cases ORDER BY id').fetchall()]
   output=[{'case':r,'record':read_record(c,r).model_dump(mode='json')} for r in cases]
   audit(c,u['id'],'export','complete_registry',detail={'records':len(output)})
  return StreamingResponse(io.BytesIO(json.dumps({'schema_version':METHOD,'records':output},ensure_ascii=False).encode()),media_type='application/json',headers={'Content-Disposition':'attachment; filename="registre_complet.json"'})
 @app.get('/registry/reports/{format}')
 def report(format:str,start:date|None=None,end:date|None=None,as_of:date|None=None,cabinet:str|None=None,category:str|None=None,provenance:str|None=None,waiting_days:int=90,u=Depends(require_roles('admin','supervisor'))):
  if format not in ('xlsx','pdf'):raise HTTPException(404,'Format indisponible')
  result=statements(start,end,as_of,cabinet,category,provenance,waiting_days,u)
  with db(write=True) as c:audit(c,u['id'],'export','operational_report',detail={'format':format,'method':METHOD,'period':result['period']})
  buf=io.BytesIO()
  if format=='xlsx':
   from openpyxl import Workbook
   w=Workbook();s=w.active;s.title='Activité';s.append(['Indicateur','Valeur'])
   for k,v in result['counts'].items():s.append([k,v])
   s.append(['Stock à la date d’observation',result['stock']['count']]);s.append(['Taux sorties / entrées',result['throughput_percent']])
   for title,key in [('Contentieux','by_category'),('Charge','by_cabinet')]:
    s=w.create_sheet(title);s.append(['Libellé','Nombre'])
    for r in result[key]:s.append([safe_cell(r['label']),r['count']])
   s=w.create_sheet('Délais');s.append(['Intervalle','Observations','Moyenne jours','Médiane jours'])
   for k,v in result['delays'].items():s.append([k,v['n'],v['mean_days'],v['median_days']])
   s=w.create_sheet('Finances');s.append(['Nature','Devise','Base de valorisation','Montant exact'])
   for r in result['financial']:s.append(list(r.values()))
   s=w.create_sheet('Qualité');s.append(['Dossier','Code','Observation'])
   for r in result['quality']['issues']:s.append([r['case_id'],r['code'],safe_cell(r['message'])])
   s=w.create_sheet('Méthode')
   for pair in [('Version',METHOD),('Période',str(result['period'])),('Filtres',str({'cabinet':cabinet,'category':category,'provenance':provenance})),('Population',result['population']),('Complétude',result['quality']['denominator']),('Prévision',str(result['forecast'])),('Avertissements',' '.join(result['warnings']))]:s.append([pair[0],safe_cell(pair[1])])
   for s in w:
    s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
    for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(70,max(18,max(len(str(c.value or '')) for c in col)+2))
   w.save(buf)
  else:
   from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle,KeepTogether
   from reportlab.lib import colors
   from reportlab.lib.styles import getSampleStyleSheet
   from reportlab.lib.pagesizes import A4
   from xml.sax.saxutils import escape
   import reportlab
   from reportlab.pdfbase import pdfmetrics
   from reportlab.pdfbase.ttfonts import TTFont
   from pathlib import Path
   fonts=Path(reportlab.__file__).parent/'fonts'
   if 'PPEFVera' not in pdfmetrics.getRegisteredFontNames():
    pdfmetrics.registerFont(TTFont('PPEFVera',str(fonts/'Vera.ttf')))
    pdfmetrics.registerFont(TTFont('PPEFVeraBold',str(fonts/'VeraBd.ttf')))
   styles=getSampleStyleSheet()
   styles['Normal'].fontName='PPEFVera';styles['Normal'].fontSize=9;styles['Normal'].leading=14
   styles['Title'].fontName='PPEFVeraBold';styles['Title'].fontSize=19;styles['Title'].leading=26
   styles['Heading2'].fontName='PPEFVeraBold';styles['Heading2'].fontSize=12;styles['Heading2'].leading=18;styles['Heading2'].keepWithNext=True
   styles['Title'].textColor=colors.HexColor('#1d4b50');styles['Heading2'].textColor=colors.HexColor('#1d4b50')
   story=[Paragraph('SISAJ / PPEF — Pilotage du greffe',styles['Title']),Paragraph('Six états · Données synthétiques · Méthode provisoire',styles['Normal']),Spacer(1,12)]
   def paragraph(text):return Paragraph(escape(str(text)),styles['Normal'])
   def section(title,headers,rows):
    section_heading=[Spacer(1,14),Paragraph(title,styles['Heading2'])]
    if not rows:story.extend(section_heading+[paragraph('Aucune observation exploitable dans ce périmètre.')]);return
    data=[[paragraph(v) for v in headers]]+[[paragraph('—' if v is None else v) for v in row] for row in rows]
    width=(A4[0]-90)/len(headers)
    table=LongTable(data,colWidths=[width]*len(headers),repeatRows=1,hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e9f1ef')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.6,colors.HexColor('#9abbb2')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8faf9')]),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
    if len(rows)<=10:story.append(KeepTogether(section_heading+[table]))
    else:story.extend(section_heading+[table])
   story.append(paragraph('Période : '+str(result['period']['start'] or 'toutes dates')+' au '+str(result['period']['end'] or 'toutes dates')+' · Observation : '+result['period']['as_of']))
   story.append(paragraph('Cabinet : '+(cabinet or 'tous')+' · Provenance : '+(provenance or 'toutes')+' · Catégorie : '+(category or 'toutes')))
   labels={'received':'Dossiers reçus','closed':'Dossiers clôturés','judged_cases':'Dossiers jugés','merits_decisions':'Décisions au fond','convicted_cases':'Dossiers avec condamnation','acquitted_cases':'Dossiers avec relaxe','mixed_decisions':'Décisions mixtes'}
   section('1. Activité générale',['Indicateur','Valeur'],[[v,result['counts'].get(k,0)] for k,v in labels.items()]+[['Stock à la date d’observation',result['stock']['count']],['Sorties / entrées (%)',result['throughput_percent']]])
   section('2. Répartition du contentieux',['Qualification source','Dossiers reçus'],[[r['label'],r['count']] for r in result['by_category']])
   section('3. Délais et attentes',['Intervalle','Nombre','Moyenne (jours)','Médiane (jours)'],[[{'receipt_to_closure':'Réception / clôture','receipt_to_hearing':'Réception / audience tenue','receipt_to_decision':'Réception / décision'}.get(k,k),v['n'],v['mean_days'],v['median_days']] for k,v in result['delays'].items()])
   section('4. Charge par cabinet',['Cabinet','Dossiers reçus'],[[r['label'],r['count']] for r in result['by_cabinet']])
   section('5. Biens et montants',['Nature','Devise','Base de valorisation','Montant exact'],[[{'offence':'Infractions','fine':'Amendes','seizure':'Saisies','confiscation':'Confiscations','compensation':'Réparations'}.get(r['kind'],r['kind']),r['currency'],CATALOGUES['valuation_basis'].get(r['basis'],r['basis']),r['amount']] for r in result['financial']])
   section('6. Qualité des données',['Indicateur','Valeur'],[['Points à examiner',result['quality']['total_issues']],['Champs renseignés',result['quality']['complete']],['Champs attendus',result['quality']['expected']],['Complétude (%)',result['quality']['completeness_percent']]])
   story.append(paragraph('Dénominateur : '+result['quality']['denominator']))
   story.extend([Spacer(1,14),Paragraph('Prévisions et méthode',styles['Heading2']),paragraph(result['forecast'].get('reason') or result['forecast'].get('warning'))])
   if result['forecast']['available']:story.append(paragraph('Estimation mensuelle : '+str(result['forecast']['monthly_estimate_next_3_months'])+' · Erreur absolue moyenne du backtest : '+str(result['forecast']['mae'])))
   for text in result['warnings']:story.extend([Spacer(1,6),paragraph(text)])
   def footer(canvas,doc):
    canvas.saveState();canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#657b80'));canvas.drawString(45,28,METHOD+' · Non officiel');canvas.drawRightString(A4[0]-45,28,str(doc.page));canvas.restoreState()
   SimpleDocTemplate(buf,pagesize=A4,leftMargin=45,rightMargin=45,topMargin=40,bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
  buf.seek(0);return StreamingResponse(buf,media_type='application/pdf' if format=='pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="ppef_six_etats.{format}"'})
 @app.get('/registry/records.xlsx')
 def complete_workbook(u=Depends(require_roles('admin','supervisor'))):
  from openpyxl import Workbook
  def flatten(o,prefix=''):
   output={}
   for k,v in o.items():
    if isinstance(v,dict):output.update(flatten(v,prefix+k+'.'))
    elif isinstance(v,list):output[prefix+k]=json.dumps(v,ensure_ascii=False)
    else:output[prefix+k]=v
   return output
  tables={k:[] for k in ('Dossiers','Personnes','Infractions','Décisions','Peines','Mesures','Résultats')}
  with db(write=True) as c:
   for case in c.execute('SELECT * FROM cases ORDER BY id').fetchall():
    case=dict(case);r=read_record(c,case).model_dump(mode='json')
    base={k:v for k,v in r.items() if k not in ('people','offences','decisions','sentences','measures')}
    tables['Dossiers'].append(flatten({**case,**base}))
    for key,title in [('people','Personnes'),('offences','Infractions'),('decisions','Décisions'),('sentences','Peines'),('measures','Mesures')]:
     for row in r[key]:
      row=dict(row);outcomes=row.pop('outcomes',[])
      tables[title].append({'case_id':case['id'],**flatten(row)})
      for o in outcomes:tables['Résultats'].append({'case_id':case['id'],'decision_reference':row['reference'],**o})
   audit(c,u['id'],'export','complete_registry',detail={'format':'xlsx','rows':len(tables['Dossiers'])})
  w=Workbook();w.remove(w.active)
  for name,rows in tables.items():
   ws=w.create_sheet(name);keys=list(dict.fromkeys(k for row in rows for k in row)) or ['case_id','Aucune donnée']
   ws.append(keys)
   for row in rows:ws.append([safe_cell(row.get(k)) for k in keys])
   ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
   for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=24
  meta=w.create_sheet('Méthode');meta.append(['Version',METHOD]);meta.append(['Population','Tous les dossiers, tous statuts et toutes dates'])
  meta.append(['Relations','case_id rattache les détails au dossier ; références locales par personne, décision et mesure'])
  meta.append(['Conservation','Original Excel complet disponible depuis le lot d’import ; aucun calcul Excel n’est assimilé à une donnée source'])
  meta.append(['Montants','Décimaux exacts en texte ; états inconnus séparés de zéro ; bases de valorisation et devises conservées'])
  buf=io.BytesIO();w.save(buf);buf.seek(0)
  return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="registre_ppef_complet.xlsx"'})
