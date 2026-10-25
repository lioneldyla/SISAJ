from copy import deepcopy
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from io import BytesIO
import json
from fastapi.testclient import TestClient
from openpyxl import Workbook,load_workbook
import pytest
from app.main import app,db
from app.registry_models import Register

PASSWORD='Synthetic-Complete-2026!'
@pytest.fixture
def context():
 with TestClient(app) as c:
  c.post('/auth/bootstrap',json={'username':'register_admin','password':PASSWORD})
  h={'Authorization':'Bearer '+c.post('/auth/login',json={'username':'register_admin','password':PASSWORD}).json()['access_token']}
  core={'rp_number':'SYNTH-REGISTER','rp_year':2026,'date_rp':'2026-01-10','provenance':'Unité fictive','category':'Qualification fictive','transnational':False}
  cid=c.post('/cases',headers=h,json=core).json()['id']
  yield c,h,cid,core

def complete():
 return {'jurisdiction_name':'Juridiction fictive','transnational_state':'undetermined','count_basis':'enumerated',
 'procedure':{'orientation':'prosecution','orientation_date':'2026-01-15','cabinet':'Cabinet fictif A','closure_date':'2026-02-01','closure_kind':'total_dismissal','first_hearing_planned':'2026-02-10','first_hearing_held':'2026-02-12','transmission_date':'2026-03-01'},
 'people':[{'reference':'P1','kind':'physical','display_name':'Personne fictive','profession':'other','profession_other':'Profession fictive','sex':'female','nationality':'Pays fictif','public_agent':'yes'},{'reference':'P2','kind':'legal','display_name':'Société fictive','activity':'finance','legal_form':'sarl'}],
 'offences':[{'reference':'I1','parquet_label':'Fictif Parquet','tribunal_label':'Fictif tribunal','offence_group':'corruption','amount':{'state':'known','value':'999.50','currency':'XOF'},'amount_basis':'alleged_damage'}],
 'decisions':[{'reference':'D1','number':'DEC-FICTIVE','decision_date':'2026-02-20','nature':'merits','result':'mixed','content':'Dispositif fictif','outcomes':[{'person_reference':'P1','offence_reference':'I1','result':'conviction'},{'person_reference':'P2','offence_reference':'I1','result':'acquittal'}]}],
 'sentences':[{'reference':'S1','decision_reference':'D1','person_reference':'P1','sentence_kind':'custody','duration_months':12,'source_text':'Douze mois fictifs'}],
 'measures':[{'reference':'M1','kind':'seizure','measure_date':'2026-01-20','object_kind':'vehicle','description':'Véhicule fictif','asset_reference':'BIEN-1','valuation_basis':'estimated','amount':{'state':'known','value':'100','currency':'XOF'}},{'reference':'M2','kind':'confiscation','decision_reference':'D1','measure_date':'2026-02-20','object_kind':'vehicle','asset_reference':'BIEN-1','disposition':'confiscated','valuation_basis':'estimated','amount':{'state':'known','value':'0','currency':'XOF'}},{'reference':'M3','kind':'fine','decision_reference':'D1','measure_date':'2026-02-20','amount':{'state':'known','value':'0.30','currency':'EUR'}},{'reference':'M4','kind':'compensation','measure_date':'2026-02-20','amount':{'state':'undetermined'}}],
 'declarations':{'convicted_total':1,'convicted_male':0,'convicted_female':1,'victim_total':2,'convicted_nationalities':'Déclaration distincte'}}

def save(c,h,cid,r,version=1):return c.put(f'/cases/{cid}/register',headers=h,json={'record':r,'expected_version':version,'reason':'Complément synthétique complet'})
def validate(c,h,cid,v=2):
 assert c.post(f'/cases/{cid}/submit?expected_version={v}',headers=h).status_code==200
 assert c.post(f'/cases/{cid}/validate?expected_version={v+1}',headers=h).status_code==200

def test_full_model_counts_states_and_round_trip(context):
 c,h,cid,_=context;r=complete();assert save(c,h,cid,r).status_code==200
 d=c.get(f'/cases/{cid}/register',headers=h).json();assert d['record']==Register.model_validate(r).model_dump(mode='json')
 core=c.get(f'/cases/{cid}',headers=h).json();assert core['transnational'] is None
 assert (core['person_physical_count'],core['person_legal_count'])==(1,1)
 assert c.put(f'/cases/{cid}/legal',headers=h,json={'expected_version':2,'reason':'Ancien éditeur'}).status_code==409
 validate(c,h,cid)
 assert save(c,h,cid,r,4).status_code==409


def test_six_statements_independent_oracle_and_exports(context):
 c,h,cid,_=context;assert save(c,h,cid,complete()).status_code==200;validate(c,h,cid)
 d=c.get('/registry/statements?start=2026-02-01&end=2026-02-28',headers=h).json()
 assert d['counts']['closed']==1 and d['counts']['judged_cases']==1 and d['counts']['mixed_decisions']==1
 assert d['counts'].get('received',0)==0 and d['throughput_percent'] is None
 assert d['stock']['count'] is None and d['instruction_stock']['count']==0
 assert d['delays']['receipt_to_closure']=={'n':1,'mean_days':22.0,'median_days':22}
 assert d['delays']['receipt_to_hearing']['median_days']==33
 assert d['delays']['receipt_to_decision']['median_days']==41
 assert d['financial']==[{'kind':'confiscation','currency':'XOF','basis':'estimated','amount':'0.00'},{'kind':'fine','currency':'EUR','basis':'undetermined','amount':'0.30'}]
 assert d['quality']['missing']['compensation_amount']==1
 assert d['sentence_months']=={'min':12,'max':12}
 assert not d['forecast']['available']
 assert c.get('/stats/legal',headers=h).json()['mixed_decisions']==1
 assert c.get('/registry/statements?as_of=2026-01-31',headers=h).json()['instruction_stock']['count']==1
 assert c.get('/registry/statements?cabinet=ABSENT',headers=h).json()['instruction_stock']['count']==0
 assert c.get('/registry/reports/pdf?start=2026-02-01&end=2026-02-28',headers=h).content.startswith(b'%PDF-')
 wb=load_workbook(BytesIO(c.get('/registry/reports/xlsx',headers=h).content));assert 'Méthode' in wb.sheetnames
 wb=load_workbook(BytesIO(c.get('/registry/records.xlsx',headers=h).content));assert set(('Personnes','Infractions','Décisions','Peines','Mesures','Résultats'))<=set(wb.sheetnames)
 rows=list(wb['Personnes'].values);assert rows[1][rows[0].index('profession_other')]=='Profession fictive'
 assert len(c.get('/registry/export.json',headers=h).json()['records'])==1

@pytest.mark.parametrize('kind', ['other_missing','unknown_money_value','negative_money','precision','foreign_decision','dates','identity_unknown_named','count_kind','sex_totals','fractional_sentence'])
def test_invalid_full_records_do_not_mutate(context,kind):
 c,h,cid,_=context;r=complete()
 if kind=='other_missing':r['people'][0]['profession_other']=None
 if kind=='unknown_money_value':r['measures'][3]['amount']['value']='0'
 if kind=='negative_money':r['measures'][0]['amount']['value']='-1'
 if kind=='precision':r['measures'][0]['amount']['value']='1.001'
 if kind=='foreign_decision':r['sentences'][0]['decision_reference']='ABSENT'
 if kind=='dates':r['procedure']['orientation_date']='2025-01-01'
 if kind=='identity_unknown_named':r['people'][0]['identity_state']='unknown'
 if kind=='count_kind':r['people'][0]['kind']='undetermined'
 if kind=='sex_totals':r['declarations']['convicted_female']=4
 if kind=='fractional_sentence':r['sentences'][0]['duration_months']=1.5
 assert save(c,h,cid,r).status_code==422
 assert c.get(f'/cases/{cid}',headers=h).json()['version']==1


def test_conditional_requirements_block_submission_not_draft(context):
 c,h,cid,_=context;r=complete();r['decisions'][0]['number']=None
 assert save(c,h,cid,r).status_code==200
 assert c.post(f'/cases/{cid}/submit?expected_version=2',headers=h).status_code==422
 r['decisions'][0]['number']='DEC-COMPLETE';assert save(c,h,cid,r,2).status_code==200;validate(c,h,cid,3)


def test_concurrent_complete_edits_have_one_winner(context):
 c,h,cid,_=context
 def update(_):
  with TestClient(app) as x:return save(x,h,cid,complete()).status_code
 with ThreadPoolExecutor(max_workers=2) as p:assert sorted(p.map(update,range(2)))==[200,409]


def csv_source(rp='SYNTH-IMPORT'):
 return ('N° RP,Date RP,Unité (OPJ) de Provenance,Dossier à caractère transnational (oui/non),Nom du prévenu,Nombre prévenus (PP),Nombre prévenus (PM),Infraction retenue par le Parquet,profession (personne publique ),Montant saisie,Objet Saisi,Date de la saisie\n'+rp+',2026-01-01,Fictif,À déterminer,Fictif,1,0,Fictif,Profession exacte,0,Véhicule,2026-02-01\n').encode()

def test_complete_import_preserves_original_and_fields(context):
 c,h,_,_=context;source=csv_source()
 d=c.post('/registry/imports/preview',headers=h,files={'file':('synthetic.csv',source,'text/csv')}).json()
 assert d['can_confirm'] and len(d['coverage'])==12
 result=c.post('/registry/imports/'+d['preview_id']+'/confirm',headers=h);assert result.status_code==200,result.text
 assert c.get('/registry/imports/'+result.json()['batch_id']+'/source',headers=h).content==source
 cid=c.get('/cases?q=SYNTH-IMPORT',headers=h).json()['items'][0]['id']
 r=c.get(f'/cases/{cid}/register',headers=h).json();assert r['record']['people'][0]['profession_other']=='Profession exacte'
 assert r['record']['measures'][0]['amount']['value']=='0.00' and r['sources'][0]['row_number']==2
 assert r['record']['transnational_state']=='undetermined'
 assert c.post('/registry/imports/'+d['preview_id']+'/confirm',headers=h).status_code==404
 assert c.post('/registry/imports/preview',headers=h,files={'file':('synthetic.csv',source)}).json()['can_confirm'] is False


def test_unmapped_field_and_legacy_loss_are_blocked(context):
 c,h,_,_=context;source=csv_source().decode().splitlines();source[0]+=',Champ inconnu';source[1]+=',valeur à préserver';source=('\n'.join(source)+'\n').encode()
 d=c.post('/registry/imports/preview',headers=h,files={'file':('synthetic.csv',source)}).json();assert not d['can_confirm'] and 'sans correspondance' in str(d['errors'])
 assert c.post('/imports/preview',headers=h,files={'file':('synthetic.csv',csv_source())}).status_code==422


def test_import_owner_and_expiry_before_consumption(context):
 c,h,_,_=context;d=c.post('/registry/imports/preview',headers=h,files={'file':('synthetic.csv',csv_source())}).json()
 c.post('/users',headers=h,json={'username':'second_agent','password':PASSWORD,'role':'collector'})
 a={'Authorization':'Bearer '+c.post('/auth/login',json={'username':'second_agent','password':PASSWORD}).json()['access_token']}
 assert c.post('/registry/imports/'+d['preview_id']+'/confirm',headers=a).status_code==403
 assert c.post('/registry/imports/'+d['preview_id']+'/confirm',headers=h).status_code==200
 for url in ('/registry/records.xlsx','/registry/reports/pdf','/registry/export.json'):assert c.get(url,headers=a).status_code==403


def test_empty_main_fallback_and_explicit_date_meaning(context):
 c,h,_,_=context;w=Workbook();s=w.active;s.title='AUDIENCES PENALES';s.append(['Titre']);s.append(['Contexte']);s.append(['N° RP','Date RP','Unité (OPJ) de Provenance'])
 s=w.create_sheet('Feuil1');s.append(['Titre']);s.append(['Contexte']);s.append(['Notice']);s.append(['N° RP','Date','Provenance','Infraction']);s.append(['SYNTH-F1','2026-01-02','Fictif','Fictif'])
 b=BytesIO();w.save(b)
 d=c.post('/registry/imports/preview',headers=h,files={'file':('synthetic.xlsx',b.getvalue())}).json()
 assert d['source_sheet']=='Feuil1' and not d['can_confirm'] and 'Date RP' in str(d['errors'])
 d=c.post('/registry/imports/preview',headers=h,data={'date_role':'rp'},files={'file':('synthetic.xlsx',b.getvalue())}).json()
 assert d['can_confirm'] and d['warnings']


def test_24_month_forecast_backtest_constant_series(context):
 c,h,cid,core=context
 with db(write=True) as conn:
  conn.execute('DELETE FROM cases WHERE id=?',(cid,))
  for i in range(30):
   y=2023+i//12;m=1+i%12
   conn.execute('INSERT INTO cases(rp_number,rp_year,date_rp,provenance,transnational,category,status,created_by,updated_by,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',(f'SYNTH-MONTH-{i}',y,f'{y}-{m:02}-01','Fictif',0,'Fictif','validated',1,1,'2026','2026'))
 d=c.get('/registry/statements?as_of=2025-07-01',headers=h).json()['forecast']
 assert d['available'] and d['mae']==0 and d['monthly_estimate_next_3_months']==1


def test_template_constants_and_formulas_do_not_create_fake_records(context):
 c,h,_,_=context;w=Workbook();s=w.active;s.title='AUDIENCES PENALES';s.append(['Titre']);s.append(['Contexte']);s.append(['N° RP','Date RP','Provenance'])
 s=w.create_sheet('Feuil1');s.append(['Titre']);s.append(['Contexte']);s.append(['Notice']);s.append(['N° RP','Date','Matière','Colonne12']);s.append([None,None,'Pénal','=IF(A5="","",1)'])
 b=BytesIO();w.save(b);d=c.post('/registry/imports/preview',headers=h,files={'file':('synthetic.xlsx',b.getvalue())}).json()
 assert d['total_rows']==0 and d['valid_rows']==0 and d['error_count']==0 and not d['can_confirm']


def test_two_source_rows_one_dossier_two_decisions(context):
 c,h,_,_=context
 source=b'rp_number,date_rp,provenance,category,transnational,N\xc2\xb0 de la d\xc3\xa9cision,Date de la d\xc3\xa9cision\nSYNTH-GROUP,2026-01-01,Fictif,Fictif,Non,D1,2026-02-01\nSYNTH-GROUP,2026-01-01,Fictif,Fictif,Non,D2,2026-03-01\n'
 d=c.post('/registry/imports/preview',headers=h,files={'file':('synthetic.csv',source)}).json()
 assert d['total_rows']==2 and d['valid_rows']==1 and d['can_confirm']
 result=c.post('/registry/imports/'+d['preview_id']+'/confirm',headers=h).json();assert result['imported_rows']==1
 cid=c.get('/cases?q=SYNTH-GROUP',headers=h).json()['items'][0]['id']
 assert len(c.get(f'/cases/{cid}/register',headers=h).json()['record']['decisions'])==2


def test_all_76_columns_are_accounted_for_and_full_source_is_reversible(context):
 from pathlib import Path
 from app.ppef_import import norm
 c,h,_,_=context
 mapping=json.loads((Path(__file__).resolve().parent.parent/'docs/ppef_column_mapping.json').read_text())
 headers=[x['column'] for x in mapping['coverage']];assert len(headers)==76
 values={norm(k):v for k,v in {'N° RP':'SYNTH-76','Date RP':'2026-01-01','Unité (OPJ) de Provenance':'OPJ fictif','Dossier à caractère transnational (oui/non)':'À déterminer','Nom du prévenu':'Personne fictive','Nombre prévenus (PP)':1,'Nombre prévenus (PM)':0,'Type de prévenus':'Personne physique','profession (personne publique )':'Fonctionnaire de l’État','Infraction retenue par le Parquet':'Qualification fictive','montant infraction':'100.50','Suite donnée':'Poursuites engagées','N° RI':'RI-FICTIF','Cabinet':'Cabinet fictif','Date de clôture ordonnance':'2026-02-01',"Nature de l'ordonnance de clôture":'Ordonnance de renvoi','Date de la saisie':'2026-01-10','Objet Saisi':'Véhicule','Montant saisie':'0','Date de la 1ère audience':'2026-02-10','Infraction retenue par le tribunal':'Qualification fictive tribunal','N° de la décision':'DEC-FICTIF','Date de la décision':'2026-03-01','Caractère décision au fond':'Oui','Contenu de la décision':'Condamnation','Peines prononcées':'Peine fictive source','Amende en fcfa':'0.10','Objet confisqué':'Véhicule','Montant confisqué':'0','Est-ce un cas de traite des personnes (OUI/NON)':'Non','Total condamné':1,'Dont condamné masculin':1,'Dont condamné féminin':0,'Total victimes':2,'Observations':'Observation intégrale fictive'}.items()}
 w=Workbook();s=w.active;s.title='AUDIENCES PENALES';s.append(['Titre']);s.append(['Contexte']);s.append(headers);s.append([values.get(norm(x)) if i<40 else '=1+1' for i,x in enumerate(headers)])
 b=BytesIO();w.save(b);source=b.getvalue()
 d=c.post('/registry/imports/preview',headers=h,files={'file':('synthetic76.xlsx',source)}).json()
 assert len(d['coverage'])==76 and d['can_confirm'],d
 result=c.post('/registry/imports/'+d['preview_id']+'/confirm',headers=h).json()
 cid=c.get('/cases?q=SYNTH-76',headers=h).json()['items'][0]['id']
 r=c.get(f'/cases/{cid}/register',headers=h).json()['record']
 assert r['source_context']['party_type_source']=='Personne physique'
 assert r['procedure']['cabinet']=='Cabinet fictif' and r['procedure']['observations']=='Observation intégrale fictive'
 assert r['source_context']['first_hearing_source_date']=='2026-02-10' and r['procedure']['first_hearing_held'] is None
 assert r['offences'][0]['amount']['value']=='100.50'
 assert r['sentences'][0]['source_text']=='Peine fictive source' and len(r['measures'])==3
 assert c.get('/registry/imports/'+result['batch_id']+'/source',headers=h).content==source


def test_creation_can_explicitly_keep_transnational_unknown(context):
 c,h,_,core=context
 data={**core,'rp_number':'SYNTH-UNKNOWN','transnational':None}
 result=c.post('/cases',headers=h,json=data);assert result.status_code==200,result.text
 cid=result.json()['id'];detail=c.get(f'/cases/{cid}',headers=h).json()
 assert detail['version']==1 and detail['transnational'] is None
 assert c.get(f'/cases/{cid}/register',headers=h).json()['record']['transnational_state']=='undetermined'


def test_new_column_after_76_is_not_misclassified_as_a_calculation(context):
 from pathlib import Path
 c,h,_,_=context
 headers=[x['column'] for x in json.loads((Path(__file__).resolve().parent.parent/'docs/ppef_column_mapping.json').read_text())['coverage']]+['Information nouvelle']
 w=Workbook();s=w.active;s.title='AUDIENCES PENALES';s.append(['Titre']);s.append(['Contexte']);s.append(headers);values=[None]*77
 values[0]='SYNTH-EXTRA';values[1]='2026-01-01';values[2]='Fictif';values[10]='Fictif';values[76]='Ne jamais ignorer'
 s.append(values);buf=BytesIO();w.save(buf)
 d=c.post('/registry/imports/preview',headers=h,files={'file':('synthetic.xlsx',buf.getvalue())}).json()
 assert not d['can_confirm'] and d['coverage'][-1]['handling']=='unmapped'


def test_csv_values_beyond_headers_are_blocked_and_preserved(context):
 c,h,_,_=context
 source=b'rp_number,date_rp,provenance,category\nSYNTH-OVERFLOW,2026-01-01,Fictif,Fictif,information sans en-tete\n'
 result=c.post('/registry/imports/preview',headers=h,files={'file':('synthetic.csv',source)}).json()
 assert not result['can_confirm']
 assert result['coverage'][-1]['column']=='COL_5' and result['coverage'][-1]['handling']=='unmapped'
 assert 'COL_5' in str(result['errors'])
 assert c.post('/registry/imports/'+result['preview_id']+'/confirm',headers=h).status_code==422
 assert not c.get('/cases?q=SYNTH-OVERFLOW',headers=h).json()['items']


@pytest.mark.parametrize('text', ['Non condamné','Annulation de la condamnation','Relaxe demandée ; décision en délibéré','Condamnation ?'])
def test_descriptive_decision_requires_human_classification(context,text):
 import csv,io
 c,h,_,_=context;buf=io.StringIO();writer=csv.writer(buf)
 writer.writerow(['rp_number','date_rp','provenance','category','N° de la décision','Date de la décision','Caractère décision au fond','Contenu de la décision'])
 writer.writerow(['SYNTH-REVIEW','2026-01-01','Fictif','Fictif','D1','2026-02-01','Oui',text]);source=buf.getvalue().encode()
 preview=c.post('/registry/imports/preview',headers=h,files={'file':('synthetic.csv',source)}).json()
 assert preview['can_confirm'] and any('résultat' in w['message'].lower() for w in preview['warnings'])
 assert c.post('/registry/imports/'+preview['preview_id']+'/confirm',headers=h).status_code==200
 cid=c.get('/cases?q=SYNTH-REVIEW',headers=h).json()['items'][0]['id']
 record=c.get(f'/cases/{cid}/register',headers=h).json()
 assert record['record']['decisions'][0]['content']==text
 assert record['record']['decisions'][0]['result']=='undetermined'
 assert any(i['code']=='DECISION_RESULT_REVIEW' for i in record['issues'])
 assert c.post(f'/cases/{cid}/submit?expected_version=2',headers=h).status_code==422
 record['record']['decisions'][0].update(result='other',result_other='Qualification vérifiée sur la minute fictive')
 assert save(c,h,cid,record['record'],2).status_code==200
 validate(c,h,cid,3)
 assert c.get('/registry/statements',headers=h).json()['counts'].get('conviction_decisions',0)==0


def test_repeated_identical_rows_keep_single_person_and_source_links(context):
 c,h,_,_=context;lines=csv_source('SYNTH-REPEATED').splitlines();source=b'\n'.join([*lines,lines[1]])+b'\n'
 preview=c.post('/registry/imports/preview',headers=h,files={'file':('synthetic.csv',source)}).json()
 assert preview['can_confirm'] and preview['valid_rows']==1 and preview['total_rows']==2
 result=c.post('/registry/imports/'+preview['preview_id']+'/confirm',headers=h).json()
 cid=c.get('/cases?q=SYNTH-REPEATED',headers=h).json()['items'][0]['id']
 record=c.get(f'/cases/{cid}/register',headers=h).json()
 assert len(record['record']['people'])==1
 assert record['record']['declarations']['unresolved_people_text'] is None
 assert [x['row_number'] for x in record['sources']]==[2,3]
 assert c.get('/registry/imports/'+result['batch_id']+'/source',headers=h).content==source


@pytest.mark.parametrize('label,expected',[('Condamnation','conviction'),('Relaxe','acquittal'),('Condamnation et relaxe','mixed'),('Sursis à statuer','stay'),('Irrecevabilité','inadmissible')])
def test_exact_result_labels_keep_their_declared_meaning(label,expected):
 from app.ppef_import import convert
 item=convert({'rp_number':'SYNTH-EXACT','date_rp':'2026-01-01','decision_content':label,'decision_number':'D1','decision_date':'2026-02-01'},2)
 assert item['record']['decisions'][0]['result']==expected
