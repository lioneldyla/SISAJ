import hashlib
import json
from copy import deepcopy
from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import app,db
from tests.test_dashboard import context,case
from tests.test_registry import complete,save,validate

def account(c,h,username,role='supervisor',profile=None,cabinets=None):
 password='Synthetic-Compliance-2026!'
 assert c.post('/users',headers=h,json={'username':username,'password':password,'role':role}).status_code==200
 with db() as conn:uid=conn.execute('SELECT id FROM users WHERE username=?',(username,)).fetchone()['id']
 if profile:assert c.put(f'/users/{uid}/data-profile',headers=h,json={'profile':profile,'cabinets':cabinets or []}).status_code==200
 return uid,{'Authorization':'Bearer '+c.post('/auth/login',json={'username':username,'password':password}).json()['access_token']}

def test_review_is_append_only_human_evidence_not_official_approval():
 with TestClient(app) as c:
  h=context(c);d=c.get('/compliance',headers=h).json()
  assert not d['officially_approved'] and d['dataset_kind']=='synthetic'
  assert all(s['latest'] is None for s in d['scenarios'])
  value={'scenario':'mixed_decision','result':'failed','expected':'Une décision mixte unique','observed':'Résultats contradictoires observés','evidence':'Fixture fictive locale, décision D1, contrôle manuel'}
  assert c.post('/compliance/reviews',headers=h,json=value).status_code==200
  value.update(result='passed',observed='Une décision mixte et deux résultats distincts')
  assert c.post('/compliance/reviews',headers=h,json=value).status_code==200
  history=c.get('/compliance/reviews',headers=h).json()['items'];assert len(history)==2
  assert [r['result'] for r in history]==['passed','failed']
  d=c.get('/compliance',headers=h).json();assert not d['officially_approved']
  assert next(s for s in d['scenarios'] if s['code']=='mixed_decision')['latest']['result']=='passed'
  assert c.post('/compliance/reviews',headers=h,json={**value,'evidence':'oui'}).status_code==422
  assert c.post('/compliance/reviews',headers=h,json={**value,'officially_approved':True}).status_code==422
  _,reader=account(c,h,'read_direction',profile='direction')
  assert c.post('/compliance/reviews',headers=reader,json=value).status_code==403
  _,scoped=account(c,h,'scoped_manager',profile='cabinet',cabinets=['A'])
  assert c.get('/compliance',headers=scoped).status_code==403
  _,clerk=account(c,h,'clerk_user',role='collector')
  assert c.get('/compliance',headers=clerk).status_code==403

def test_snapshot_preserves_values_filters_hash_and_original_rows_after_changes():
 with TestClient(app) as c:
  h=context(c);a=case(c,h,'SYNTH-A','A');case(c,h,'SYNTH-B','B')
  r=c.post('/dashboard/snapshots',headers=h,json={'title':'Rapport fictif A','filters':{'cabinet':['A'],'as_of':'2026-04-01'}})
  assert r.status_code==200,r.text
  sid=r.json()['id'];raw=c.get(f'/dashboard/snapshots/{sid}/export/json',headers=h)
  assert hashlib.sha256(raw.content).hexdigest()==r.json()['sha256']
  d=raw.json();assert d['dataset_kind']=='synthetic' and not d['officially_approved']
  assert d['analysis']['cohort_size']==1 and d['analysis']['rows'][0]['id']==a
  assert d['analysis']['filters']['cabinet']==['A']
  record=c.get(f'/cases/{a}/register',headers=h).json()['record'];record['procedure']['cabinet']='C'
  assert c.put(f'/cases/{a}/register',headers=h,json={'expected_version':2,'reason':'Affectation synthétique corrigée','record':record}).status_code==200
  assert c.get(f'/dashboard/snapshots/{sid}/export/json',headers=h).content==raw.content
  export=c.get(f'/dashboard/snapshots/{sid}/export/csv',headers=h).text
  assert 'Fictives / test' in export and 'SYNTH-A' in export and 'SYNTH-B' not in export
  assert 'Non calculable' in export
  assert c.post('/dashboard/snapshots',headers=h,json={'title':'Invalide','filters':{'start':'2026-04-01','end':'2026-01-01'}}).status_code==422
  assert c.post('/dashboard/snapshots',headers=h,json={'title':'Invalide','filters':{'limit':0}}).status_code==422

def test_snapshot_scope_checked_for_entire_cohort_and_after_reassignment():
 with TestClient(app) as c:
  h=context(c);case(c,h,'SYNTH-A','A');case(c,h,'SYNTH-B','B')
  uid,scoped=account(c,h,'cabinet_manager',profile='cabinet',cabinets=['A'])
  r=c.post('/dashboard/snapshots',headers=scoped,json={'title':'Rapport cabinet A','filters':{'metric':'returned'}})
  assert r.status_code==200,r.text
  sid=r.json()['id'];d=c.get(f'/dashboard/snapshots/{sid}',headers=scoped).json()
  assert len(d['snapshot']['authorized_case_ids'])==1 and d['snapshot']['analysis']['rows']==[]
  _,other=account(c,h,'other_manager',profile='cabinet',cabinets=['A'])
  assert c.get(f'/dashboard/snapshots/{sid}',headers=other).status_code==404
  assert c.put(f'/users/{uid}/data-profile',headers=h,json={'profile':'cabinet','cabinets':['B']}).status_code==200
  assert c.get(f'/dashboard/snapshots/{sid}',headers=scoped).status_code==403
  assert c.get('/dashboard/snapshots',headers=scoped).json()['items']==[]

def test_snapshot_tampering_is_detected_before_export():
 with TestClient(app) as c:
  h=context(c);r=c.post('/dashboard/snapshots',headers=h,json={'title':'Rapport intégrité'}).json()
  with db(write=True) as conn:conn.execute('UPDATE analysis_snapshots SET payload=? WHERE id=?',('{}',r['id']))
  assert c.get('/dashboard/snapshots/'+r['id'],headers=h).status_code==409

def test_successive_hearings_survive_save_export_and_do_not_invent_final_stock():
 with TestClient(app) as c:
  h=context(c);cid=case(c,h,'SYNTH-A','A');r=complete()
  r['hearings']=[{'reference':'A1','planned_date':'2026-02-10','held_date':'2026-02-12','status':'adjourned','next_date':'2026-03-10','reason':'Renvoi fictif contradictoire','source_reference':'PV-1'}, {'reference':'A2','planned_date':'2026-03-10','status':'planned','source_reference':'ROLE-2'}]
  assert save(c,h,cid,r,2).status_code==200
  d=c.get('/dashboard/analysis?date_basis=hearing&start=2026-02-12&end=2026-02-12',headers=h).json()
  assert d['cohort_size']==1 and d['hearing_summary']['by_status']=={'adjourned':1}
  assert 'renvoyée' in d['rows'][0]['last_event']
  validate(c,h,cid,3)
  s=c.get('/registry/statements',headers=h).json()
  assert s['stock']['available'] is False and s['stock']['count'] is None
  assert s['instruction_stock']['count']==0 and s['throughput_percent'] is None
  wb=load_workbook(BytesIO(c.get('/registry/records.xlsx',headers=h).content))
  assert wb['Audiences'].max_row==3
  report=load_workbook(BytesIO(c.get('/registry/reports/xlsx',headers=h).content))
  rows=list(report['Activité'].values)
  assert ('Stock judiciaire final','Non calculable') in rows

@pytest.mark.parametrize('change',[{'status':'held','planned_date':'2026-02-01'}, {'status':'adjourned','held_date':'2026-02-01','next_date':'2026-01-31','reason':'Renvoi fictif'}, {'status':'cancelled','planned_date':'2026-02-01'}, {'status':'planned','planned_date':'2025-01-01'}])
def test_invalid_hearings_do_not_mutate_case(change):
 with TestClient(app) as c:
  h=context(c);cid=case(c,h,'SYNTH-A','A');r=complete();r['hearings']=[{'reference':'A1',**change}]
  assert save(c,h,cid,r,2).status_code==422
  assert c.get(f'/cases/{cid}',headers=h).json()['version']==2

@pytest.mark.parametrize('kind',['absent','unknown','conflict'])
def test_incomplete_decision_is_saved_as_draft_but_submission_is_blocked(kind):
 with TestClient(app) as c:
  h=context(c);cid=case(c,h,'SYNTH-A','A');r=complete()
  if kind=='absent':r['decisions'][0]['outcomes']=[]
  elif kind=='unknown':r['decisions'][0]['outcomes'][0]['result']='unknown'
  else:r['decisions'][0]['result']='conviction'
  result=save(c,h,cid,r,2);assert result.status_code==200
  assert any(i['code'] in ('RESULT_CONFLICT','INDIVIDUAL_RESULT') for i in result.json()['issues'])
  assert c.post(f'/cases/{cid}/submit?expected_version=3',headers=h).status_code==422
  assert c.get(f'/cases/{cid}',headers=h).json()['status']=='draft'
