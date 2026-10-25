"""Browser API and dual-storage regression tests, using synthetic records only."""
from io import BytesIO
import json
import os
import subprocess
import sys

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import app, db, DATABASE_URL

PASSWORD = 'Synthetic-Platform-2026!'

@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


def auth(c):
    assert c.post('/auth/bootstrap',json={'username':'admin_test','password':PASSWORD}).status_code==200
    r=c.post('/auth/login',json={'username':'admin_test','password':PASSWORD})
    assert r.status_code==200
    return {'Authorization':'Bearer '+r.json()['access_token']}


def record(**changes):
    return dict(rp_number='SYNTH-RP-001',rp_year=2026,date_rp='2026-01-01',provenance='Unité fictive',category='Catégorie fictive',transnational=False,**changes)


def test_cookie_csrf_logout_and_headers(client):
    auth(client)
    assert client.get('/auth/me').status_code==200
    assert client.post('/cases',json=record()).status_code==403
    assert client.post('/cases',headers={'X-SISA-CSRF':'1'},json=record()).status_code==200
    assert client.post('/auth/logout',headers={'X-SISA-CSRF':'1'}).status_code==200
    assert client.get('/cases').status_code==401
    response=client.get('/')
    assert 'script-src' in response.headers['content-security-policy']
    assert response.headers['x-frame-options']=='DENY'
    assert client.post('/auth/login',headers={'Origin':'https://foreign.invalid'},json={'username':'x','password':'x'}).status_code==403


def test_protected_bootstrap(client,monkeypatch):
    monkeypatch.setenv('SISA_BOOTSTRAP_KEY','synthetic-initialization-key')
    body={'username':'admin_test','password':PASSWORD}
    assert client.post('/auth/bootstrap',json=body).status_code==403
    assert client.post('/auth/bootstrap',headers={'X-SISA-Bootstrap-Key':'synthetic-initialization-key'},json=body).status_code==200


def test_login_throttling(client):
    for _ in range(5):
        assert client.post('/auth/login',json={'username':'unknown','password':'incorrect'}).status_code==401
    assert client.post('/auth/login',json={'username':'unknown','password':'incorrect'}).status_code==429


def test_roles_and_revocation(client):
    admin=auth(client)
    assert client.post('/users',headers=admin,json={'username':'collector_test','password':PASSWORD,'role':'collector'}).status_code==200
    r=client.post('/auth/login',json={'username':'collector_test','password':PASSWORD})
    agent={'Authorization':'Bearer '+r.json()['access_token']}
    for url in ['/audit','/users','/exports/cases.csv','/exports/cases.xlsx','/reports/summary.pdf','/intelligence/anomalies','/admin/operations']:
        assert client.get(url,headers=agent).status_code==403,url
    assert client.post('/admin/backup',headers=agent).status_code==403
    cid=client.post('/cases',headers=agent,json=record()).json()['id']
    assert client.post(f'/cases/{cid}/submit',headers=agent).status_code==200
    assert client.post(f'/cases/{cid}/validate',headers=agent).status_code==403
    assert client.post(f'/cases/{cid}/return',headers=agent,json={'reason':'synthetic refusal'}).status_code==403
    user=next(u for u in client.get('/users',headers=admin).json()['items'] if u['role']=='collector')
    assert client.patch('/users/'+str(user['id']),headers=admin,json={'active':False}).status_code==200
    assert client.get('/cases',headers=agent).status_code==401


def test_invalid_data_conflict_and_version(client):
    h=auth(client);data=record()
    a=client.post('/cases',headers=h,json=data).json()['id']
    b=client.post('/cases',headers=h,json={**data,'rp_number':'SYNTH-RP-002'}).json()['id']
    assert client.patch(f'/cases/{b}?reason=synthetic',headers=h,json=data).status_code==409
    assert client.get(f'/cases/{b}',headers=h).json()['rp_number']=='SYNTH-RP-002'
    assert client.patch(f'/cases/{a}?reason=synthetic&expected_version=99',headers=h,json=data).status_code==409
    for field in ['rp_number','category','provenance']:
        assert client.post('/cases',headers=h,json={**data,field:'   '}).status_code==422
    for val in ['NaN','Infinity','-1']:
        assert client.post(f'/cases/{a}/timeline',headers=h,json={'event_type':'Mesure fictive','fine_amount':val}).status_code==422
    assert client.post(f'/cases/{a}/timeline',headers=h,json={'event_type':'Mesure fictive','fine_amount':'100.01'}).status_code==200
    assert client.post(f'/cases/{a}/submit',headers=h).status_code==200
    assert client.post(f'/cases/{a}/validate',headers=h).status_code==200
    assert client.post(f'/cases/{a}/timeline',headers=h,json={'event_type':'Audience fictive'}).status_code==409


@pytest.mark.parametrize('field,value',[('person_physical_count','1.9'),('person_legal_count','NaN'),('rp_year','2026.9')])
def test_fractional_import_blocked(client,field,value):
    h=auth(client);data={**record(),field:value}
    content=','.join(data)+'\n'+','.join(str(v) for v in data.values())+'\n'
    r=client.post('/imports/preview',headers=h,files={'file':('synthetic.csv',content.encode(),'text/csv')})
    assert r.status_code==200
    assert not r.json()['can_confirm']
    assert client.post('/imports/'+r.json()['preview_id']+'/confirm',headers=h).status_code==422
    assert client.get('/cases',headers=h).json()['total']==0


def test_exports_dashboard_and_filter(client):
    h=auth(client);data={**record(),'category':'=1+1'}
    cid=client.post('/cases',headers=h,json=data).json()['id']
    assert client.get('/dashboard',headers=h).json()['total']==1
    assert client.get('/cases?start=2026-02-01',headers=h).json()['total']==0
    assert client.get('/cases?start=2027-01-01&end=2026-01-01',headers=h).status_code==422
    csv=client.get('/exports/cases.csv',headers=h)
    assert "'=1+1" in csv.text
    xlsx=client.get('/exports/cases.xlsx',headers=h)
    wb=load_workbook(BytesIO(xlsx.content),data_only=False)
    assert wb['Dossiers']['F2'].value=="'=1+1"
    assert 'Méthode' in wb.sheetnames
    assert client.get('/reports/summary.pdf',headers=h).content.startswith(b'%PDF-')
    assert client.get('/ready').json()['storage']==('postgresql' if DATABASE_URL else 'sqlite')
    assert client.get('/cases/'+str(cid),headers=h).json()['category']=='=1+1'


def test_preview_ownership_and_replay_on_both_backends(client):
    owner=auth(client)
    client.post('/users',headers=owner,json={'username':'other_test','password':PASSWORD,'role':'collector'})
    other={'Authorization':'Bearer '+client.post('/auth/login',json={'username':'other_test','password':PASSWORD}).json()['access_token']}
    content=b'rp_number,rp_year,date_rp,provenance,category,transnational\nSYNTH-IMP,2026,2026-02-01,OPJ fictif,Categorie fictive,Non\n'
    r=client.post('/imports/preview',headers=owner,files={'file':('synthetic.csv',content,'text/csv')})
    pid=r.json()['preview_id']
    assert client.post(f'/imports/{pid}/confirm',headers=other).status_code==403
    assert client.post(f'/imports/{pid}/confirm',headers=owner).json()['imported_rows']==1
    assert client.post(f'/imports/{pid}/confirm',headers=owner).status_code==404
    assert client.get('/imports',headers=owner).json()['items'][0]['imported_count']==1


def test_render_refuses_sqlite_fallback():
    env={**os.environ,'SISA_DEPLOYMENT':'render','DATABASE_URL':''}
    r=subprocess.run([sys.executable,'-c','import app.main'],env=env,capture_output=True,text=True)
    assert r.returncode!=0
    assert 'aucun repli SQLite' in r.stderr


def test_password_change_revokes_sessions(client):
    h=auth(client)
    assert client.post('/auth/password',headers=h,json={'current_password':PASSWORD,'new_password':'Another-Synthetic-Password-2026!'}).status_code==200
    assert client.get('/cases',headers=h).status_code==401


def test_return_correction_and_statistics_population(client):
    h=auth(client);data=record()
    cid=client.post('/cases',headers=h,json=data).json()['id']
    assert client.get('/stats',headers=h).json()['received_cases']==0
    assert client.post(f'/cases/{cid}/submit',headers=h).status_code==200
    assert client.post(f'/cases/{cid}/return',headers=h,json={'reason':'     '}).status_code==422
    assert client.post(f'/cases/{cid}/return',headers=h,json={'reason':'Source fictive à vérifier'}).status_code==200
    assert client.get(f'/cases/{cid}',headers=h).json()['status']=='returned'
    assert client.post(f'/cases/{cid}/submit',headers=h).status_code==200
    assert client.post(f'/cases/{cid}/validate',headers=h).status_code==200
    assert client.get('/stats',headers=h).json()['received_cases']==1
    changed={**data,'person_physical_count':0}
    assert client.post(f'/cases/{cid}/corrections',headers=h,json={'case':changed,'reason':'Correction synthetique','expected_version':5}).status_code==200
    assert client.get('/stats',headers=h).json()['received_cases']==0
    detail=client.get(f'/cases/{cid}',headers=h).json()
    assert detail['status']=='draft' and detail['version']==6 and detail['person_physical_count']==0
    trace=next(r for r in client.get('/audit',headers=h).json()['items'] if r['action']=='correct')
    content=json.loads(trace['detail_json'])
    assert content['before']['status']=='validated'
    assert content['after']['person_physical_count']==0
    assert content['reason']=='Correction synthetique'
